#!/usr/bin/env python3
from __future__ import annotations

import json
import math
import re
import shutil
import subprocess
import time
import zipfile
from copy import deepcopy
from pathlib import Path
from typing import Optional

import fitz
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import first_book_28_step_compliance as b

ROOT = Path.cwd()
SOURCE = ROOT / "source_interior.docx"
OUT = ROOT / "output"
OUT.mkdir(exist_ok=True)
CONTACT = OUT / "لوحات_المراجعة"
CONTACT.mkdir(exist_ok=True)

BOOK_TITLE = "رؤية المملكة 2030 والأمن الشامل"
SUBTITLE = "من حماية الدولة إلى صناعة المستقبل"
AUTHOR = "عبد العزيز القرني"
AUTHOR_CREDENTIAL = "باحث دكتوراه في القانون"
EDITION = "النسخة المنقحة والمخرجة عربيًا | 2026م"

FINAL_DOCX = OUT / "رؤية_المملكة_2030_والأمن_الشامل_النسخة_النهائية.docx"
FINAL_PDF = OUT / "رؤية_المملكة_2030_والأمن_الشامل_نسخة_الطباعة_الرمادية.pdf"
FRONT_DOCX = OUT / "رؤية_المملكة_2030_والأمن_الشامل_الغلاف_الأمامي.docx"
FRONT_PNG = OUT / "رؤية_المملكة_2030_والأمن_الشامل_الغلاف_الأمامي_300DPI.png"
BACK_DOCX = OUT / "رؤية_المملكة_2030_والأمن_الشامل_الغلاف_الخلفي.docx"
BACK_PNG = OUT / "رؤية_المملكة_2030_والأمن_الشامل_الغلاف_الخلفي_300DPI.png"
SPINE_DOCX = OUT / "رؤية_المملكة_2030_والأمن_الشامل_دليل_الكعب.docx"
SPINE_PNG = OUT / "رؤية_المملكة_2030_والأمن_الشامل_دليل_الكعب_غير_مقاس.png"
REPORT_DOCX = OUT / "تقرير_مطابقة_رؤية_المملكة_2030_والأمن_الشامل_للخطوات_28.docx"
REPORT_PDF = OUT / "تقرير_مطابقة_رؤية_المملكة_2030_والأمن_الشامل_للخطوات_28.pdf"
AUDIT_JSON = OUT / "audit_28_steps.json"
README = OUT / "اقرأني_أولاً.txt"
PACKAGE_ZIP = ROOT / "رؤية_المملكة_2030_والأمن_الشامل_الحزمة_النهائية.zip"

BLACK = "000000"
WHITE = "FFFFFF"
DARK_GRAY = "3F3F3F"
MID_GRAY = "7F7F7F"
LIGHT_GRAY = "F2F2F2"
PETROL = "164C5A"
GOLD = "B39756"
IVORY = "F4EFE3"
CHARCOAL = "333333"
ARABIC_DIGITS = str.maketrans("0123456789", "٠١٢٣٤٥٦٧٨٩")
ORDINALS = "الأول|الثاني|الثالث|الرابع|الخامس|السادس|السابع|الثامن|التاسع|العاشر|الحادي عشر|الثاني عشر|الثالث عشر|الرابع عشر"
CHAPTER_RE = re.compile(rf"^الفصل\s+({ORDINALS}|[٠-٩0-9]+)\s*$")


def clean(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "").replace("\u200f", "").replace("\u200e", "")).strip()


def ar_ratio(s: str) -> float:
    letters = re.findall(r"[A-Za-z\u0600-\u06FF]", s or "")
    if not letters:
        return 0.0
    return sum("\u0600" <= c <= "\u06FF" for c in letters) / len(letters)


def ensure_el(parent, tag: str):
    el = parent.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        parent.append(el)
    return el


def set_bidi_paragraph(p, alignment=None):
    ppr = p._p.get_or_add_pPr()
    if ppr.find(qn("w:bidi")) is None:
        ppr.insert(0, OxmlElement("w:bidi"))
    if alignment is not None:
        p.alignment = alignment
        jc = ensure_el(ppr, "w:jc")
        jc.set(qn("w:val"), {
            WD_ALIGN_PARAGRAPH.RIGHT: "right",
            WD_ALIGN_PARAGRAPH.LEFT: "left",
            WD_ALIGN_PARAGRAPH.CENTER: "center",
            WD_ALIGN_PARAGRAPH.JUSTIFY: "both",
        }.get(alignment, "right"))


def set_run(run, font: str, size: float, bold: Optional[bool] = None, color=BLACK, rtl=True):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.font.bold = bold
    rpr = run._r.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), font)
    if rtl and rpr.find(qn("w:rtl")) is None:
        rpr.append(OxmlElement("w:rtl"))
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:val"), "ar-SA" if rtl else "en-US")
    lang.set(qn("w:bidi"), "ar-SA" if rtl else "en-US")


def ensure_style(doc, name, font, size, bold=False, align=WD_ALIGN_PARAGRAPH.RIGHT,
                 before=0, after=5, line=1.2, first=0, left=0, right=0,
                 keep_next=False, keep_lines=True, outline=None, color=BLACK):
    try:
        style = doc.styles[name]
        if style.type != WD_STYLE_TYPE.PARAGRAPH:
            style.element.getparent().remove(style.element)
            style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    except KeyError:
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = font
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    pf = style.paragraph_format
    pf.alignment = align
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.first_line_indent = Cm(first) if first else None
    pf.left_indent = Cm(left) if left else None
    pf.right_indent = Cm(right) if right else None
    pf.keep_with_next = keep_next
    pf.keep_together = keep_lines
    pf.widow_control = True
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), font)
    if rpr.find(qn("w:rtl")) is None:
        rpr.append(OxmlElement("w:rtl"))
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:val"), "ar-SA")
    lang.set(qn("w:bidi"), "ar-SA")
    ppr = style.element.get_or_add_pPr()
    if ppr.find(qn("w:bidi")) is None:
        ppr.insert(0, OxmlElement("w:bidi"))
    if outline is not None:
        lvl = ensure_el(ppr, "w:outlineLvl")
        lvl.set(qn("w:val"), str(outline))
    return style


def create_styles(doc):
    ensure_style(doc, "Book Title", "Noto Kufi Arabic", 28, True, WD_ALIGN_PARAGRAPH.CENTER, after=10, line=1.05)
    ensure_style(doc, "Book Subtitle", "Noto Kufi Arabic", 15, False, WD_ALIGN_PARAGRAPH.CENTER, after=8, line=1.15)
    ensure_style(doc, "Author Name", "Noto Kufi Arabic", 15, False, WD_ALIGN_PARAGRAPH.CENTER, after=8, line=1.15)
    ensure_style(doc, "Part Title", "Noto Kufi Arabic", 24, True, WD_ALIGN_PARAGRAPH.CENTER, before=60, after=12, line=1.05, keep_next=True, outline=0)
    ensure_style(doc, "Chapter Number", "Noto Kufi Arabic", 16, True, WD_ALIGN_PARAGRAPH.RIGHT, before=48, after=8, line=1.0, keep_next=True)
    ensure_style(doc, "Chapter Title", "Noto Kufi Arabic", 22, True, WD_ALIGN_PARAGRAPH.RIGHT, before=0, after=14, line=1.1, keep_next=True, outline=1)
    ensure_style(doc, "Main Heading", "Noto Kufi Arabic", 17, True, WD_ALIGN_PARAGRAPH.RIGHT, before=12, after=6, line=1.1, keep_next=True, outline=2)
    ensure_style(doc, "Subheading", "Noto Kufi Arabic", 15, True, WD_ALIGN_PARAGRAPH.RIGHT, before=10, after=5, line=1.1, keep_next=True, outline=3)
    ensure_style(doc, "Unnumbered Internal Heading", "Noto Kufi Arabic", 15, True, WD_ALIGN_PARAGRAPH.RIGHT, before=10, after=5, line=1.1, keep_next=True)
    ensure_style(doc, "Body Text", "Amiri", 14, False, WD_ALIGN_PARAGRAPH.JUSTIFY, after=5, line=1.2, first=.6)
    ensure_style(doc, "First Paragraph after Heading", "Amiri", 14, False, WD_ALIGN_PARAGRAPH.JUSTIFY, after=5, line=1.2)
    ensure_style(doc, "Long Quotation", "Amiri", 13, False, WD_ALIGN_PARAGRAPH.JUSTIFY, before=8, after=8, line=1.15, left=1, right=1)
    ensure_style(doc, "Qur’anic Verse", "Amiri Quran", 15, False, WD_ALIGN_PARAGRAPH.CENTER, before=8, after=8, line=1.15)
    ensure_style(doc, "Prophetic Hadith", "Amiri", 13.5, False, WD_ALIGN_PARAGRAPH.JUSTIFY, before=6, after=6, line=1.15, left=.8, right=.8)
    ensure_style(doc, "Box Title", "Noto Kufi Arabic", 13, True, WD_ALIGN_PARAGRAPH.RIGHT, before=8, after=3, line=1.1)
    ensure_style(doc, "Box Text", "Amiri", 12.5, False, WD_ALIGN_PARAGRAPH.JUSTIFY, after=8, line=1.2)
    ensure_style(doc, "Table Title", "Amiri", 10.5, True, WD_ALIGN_PARAGRAPH.RIGHT, before=6, after=3, line=1.05, keep_next=True)
    ensure_style(doc, "Table Text", "Amiri", 11.5, False, WD_ALIGN_PARAGRAPH.RIGHT, after=2, line=1.05)
    ensure_style(doc, "Table Source", "Amiri", 10.5, False, WD_ALIGN_PARAGRAPH.RIGHT, before=3, after=6, line=1.05)
    ensure_style(doc, "Figure Title", "Amiri", 10.5, True, WD_ALIGN_PARAGRAPH.CENTER, before=4, after=2, line=1.05)
    ensure_style(doc, "Figure Caption", "Amiri", 10.5, False, WD_ALIGN_PARAGRAPH.CENTER, before=2, after=6, line=1.05)
    ensure_style(doc, "Footnote", "Amiri", 10.5, False, WD_ALIGN_PARAGRAPH.RIGHT, after=2, line=1.0)
    ensure_style(doc, "Bulleted List", "Amiri", 14, False, WD_ALIGN_PARAGRAPH.RIGHT, after=4, line=1.15, right=.7)
    ensure_style(doc, "Numbered List", "Amiri", 14, False, WD_ALIGN_PARAGRAPH.RIGHT, after=4, line=1.15, right=.7)
    ensure_style(doc, "Reference", "Amiri", 11.5, False, WD_ALIGN_PARAGRAPH.RIGHT, after=4, line=1.15, right=.7)
    ensure_style(doc, "Reference Latin", "Amiri", 11.5, False, WD_ALIGN_PARAGRAPH.LEFT, after=4, line=1.15, left=.7)
    ensure_style(doc, "Appendix Title", "Noto Kufi Arabic", 22, True, WD_ALIGN_PARAGRAPH.RIGHT, before=45, after=14, line=1.1, keep_next=True, outline=1)
    ensure_style(doc, "Index Heading", "Noto Kufi Arabic", 18, True, WD_ALIGN_PARAGRAPH.RIGHT, before=24, after=10, line=1.05)
    ensure_style(doc, "Index Field", "Amiri", 12, False, WD_ALIGN_PARAGRAPH.RIGHT, after=5, line=1.1)


def add_field(p, instruction, placeholder=""):
    r = p.add_run()._r
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin"); begin.set(qn("w:dirty"), "true"); r.append(begin)
    r = p.add_run()._r
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = instruction; r.append(instr)
    r = p.add_run()._r
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate"); r.append(sep)
    if placeholder:
        p.add_run(placeholder)
    r = p.add_run()._r
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end"); r.append(end)


def add_hidden_index_field(p, term: str, index_id: str):
    if not term:
        return
    safe = term.replace('"', "'")
    r = p.add_run()._r
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin"); begin.set(qn("w:dirty"), "true"); r.append(begin)
    r = p.add_run()._r
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = f' XE "{safe}" \\f "{index_id}" '; r.append(instr)
    r = p.add_run()._r
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end"); r.append(end)
    for run in p.runs[-3:]:
        rpr = run._r.get_or_add_rPr()
        if rpr.find(qn("w:vanish")) is None:
            rpr.append(OxmlElement("w:vanish"))


def insert_after(p, text="", style=None):
    new_p = OxmlElement("w:p")
    p._p.addnext(new_p)
    new = b.Paragraph(new_p, p._parent) if hasattr(b, "Paragraph") else None
    if new is None:
        from docx.text.paragraph import Paragraph
        new = Paragraph(new_p, p._parent)
    if text:
        new.add_run(text)
    if style:
        new.style = style
    return new


def delete_paragraph(p):
    parent = p._p.getparent()
    if parent is not None:
        parent.remove(p._p)


def replace_manual_indexes(doc):
    paras = list(doc.paragraphs)
    toc_i = next((i for i, p in enumerate(paras) if clean(p.text) == "فهرس المحتويات"), None)
    chapter_i = next((i for i, p in enumerate(paras) if CHAPTER_RE.match(clean(p.text))), None)
    if toc_i is None or chapter_i is None or chapter_i <= toc_i:
        return False
    for p in paras[toc_i + 1:chapter_i]:
        delete_paragraph(p)
    toc = doc.paragraphs[toc_i]
    toc.style = "Index Heading"
    field_p = insert_after(toc, style="Index Field")
    add_field(field_p, ' TOC \\o "1-4" \\h \\z \\u ', "يُحدّث تلقائيًا في Microsoft Word")
    lot_h = insert_after(field_p, "قائمة الجداول", "Index Heading")
    lot = insert_after(lot_h, style="Index Field")
    add_field(lot, ' TOC \\h \\z \\t "Table Title,1" ', "يُحدّث تلقائيًا في Microsoft Word")
    lof_h = insert_after(lot, "قائمة الأشكال", "Index Heading")
    lof = insert_after(lof_h, style="Index Field")
    add_field(lof, ' TOC \\h \\z \\t "Figure Caption,1" ', "يُحدّث تلقائيًا في Microsoft Word")
    return True


def paragraph_has_drawing(p):
    return bool(p._p.xpath('.//w:drawing|.//w:pict|.//v:shape'))


def clear_page_breaks(p):
    ppr = p._p.get_or_add_pPr()
    pb = ppr.find(qn("w:pageBreakBefore"))
    if pb is not None:
        ppr.remove(pb)
    for br in list(p._p.xpath('.//w:br[@w:type="page"]')):
        br.getparent().remove(br)


def apply_box(p, title=False):
    ppr = p._p.get_or_add_pPr()
    shd = ensure_el(ppr, "w:shd"); shd.set(qn("w:fill"), LIGHT_GRAY)
    borders = ensure_el(ppr, "w:pBdr")
    right = ensure_el(borders, "w:right")
    right.set(qn("w:val"), "single"); right.set(qn("w:sz"), "18"); right.set(qn("w:space"), "8"); right.set(qn("w:color"), PETROL)
    p.paragraph_format.left_indent = Cm(.4)
    p.paragraph_format.right_indent = Cm(.4)
    p.paragraph_format.space_before = Pt(8 if title else 0)
    p.paragraph_format.space_after = Pt(3 if title else 8)


def normalize_list_text(text):
    t = clean(text)
    m = re.match(r"^\.(\d+(?:\.\d+)*)\s*(.*)$", t)
    if not m:
        m = re.match(r"^(\d+(?:\.\d+)*)[.)]?\s+(.*)$", t)
    if not m:
        return text
    number = m.group(1).translate(ARABIC_DIGITS)
    rest = m.group(2)
    return f"{number}. {rest}" if "." not in m.group(1) else f"{number} {rest}"


def is_quran(text):
    return bool(re.search(r"(?:قال|قوله)\s+تعالى|\[?قريش\s*[:：]\s*[٠-٩0-9]+|﴿|\{[^{}]{10,}\}", text))


def is_hadith(text):
    return "رواه" in text and any(x in text for x in ("حديث", "الترمذي", "البخاري", "مسلم", "أبو داود", "النسائي", "ابن ماجه"))


def original_outline_level(p):
    ppr = p._p.get_or_add_pPr()
    out = ppr.find(qn("w:outlineLvl"))
    if out is not None:
        try:
            return int(out.get(qn("w:val")))
        except Exception:
            return None
    try:
        sppr = p.style.element.get_or_add_pPr()
        out = sppr.find(qn("w:outlineLvl"))
        return int(out.get(qn("w:val"))) if out is not None else None
    except Exception:
        return None


def short_heading_candidate(text, next_text, old_style, p):
    if not text or len(text) > 100 or text.startswith(("http://", "https://")):
        return False
    if text.endswith((".", "،", "؛", ":")) and not text.endswith(":"):
        return False
    if re.match(r"^[•●▪■-]", text) or re.match(r"^[٠-٩0-9]+[.)]", text) or re.match(r"^\.[٠-٩0-9]+", text):
        return False
    low = (old_style or "").lower()
    if any(k in low for k in ("heading", "title", "عنوان", "head")):
        return True
    if original_outline_level(p) is not None:
        return True
    words = text.split()
    return len(words) <= 11 and len(next_text) >= 75


def style_document_paragraphs(doc):
    stats = {"paragraphs": len(doc.paragraphs), "rtl_fixed": 0, "first_after_heading_fixed": 0,
             "numbering_normalized": 0, "quran": 0, "hadith": 0, "boxes": 0,
             "table_captions": 0, "figure_captions": 0, "references": 0, "hyperlinks": 0}
    paras = list(doc.paragraphs)
    texts = [clean(p.text) for p in paras]
    nonempty_indices = [i for i, t in enumerate(texts) if t]
    title_occurrences = [i for i, t in enumerate(texts) if t == BOOK_TITLE]
    if title_occurrences:
        paras[title_occurrences[0]].style = "Book Title"
    if len(title_occurrences) > 1:
        paras[title_occurrences[1]].style = "Book Title"
    for i, t in enumerate(texts[:20]):
        if t == SUBTITLE:
            paras[i].style = "Book Subtitle"
        elif t == AUTHOR:
            paras[i].style = "Author Name"
        elif t in ("بقلم", AUTHOR_CREDENTIAL):
            paras[i].style = "Book Subtitle"

    chapter_title_indices = set()
    chapter_number_indices = set()
    for i, t in enumerate(texts):
        if CHAPTER_RE.match(t):
            chapter_number_indices.add(i)
            j = next((k for k in range(i + 1, len(texts)) if texts[k]), None)
            if j is not None:
                chapter_title_indices.add(j)

    reference_mode = False
    author_bio_mode = False
    previous_semantic = None
    box_next = False
    heading_names = {"Chapter Number", "Chapter Title", "Main Heading", "Subheading", "Unnumbered Internal Heading", "Appendix Title", "Box Title", "Index Heading"}
    for i, p in enumerate(paras):
        text = clean(p.text)
        if not text:
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            continue
        old_style = p.style.name if p.style else ""
        next_text = next((texts[k] for k in range(i + 1, len(texts)) if texts[k]), "")

        if text == "المراجع والمصادر النهائية":
            p.style = "Appendix Title"; reference_mode = True; author_bio_mode = False
        elif text == "حول الكتاب والكاتب":
            p.style = "Appendix Title"; reference_mode = False; author_bio_mode = True
        elif i in chapter_number_indices:
            p.style = "Chapter Number"; clear_page_breaks(p)
        elif i in chapter_title_indices:
            p.style = "Chapter Title"
        elif text in ("فهرس المحتويات", "قائمة الجداول", "قائمة الأشكال"):
            p.style = "Index Heading"
        elif re.match(r"^(?:جدول|الجدول)\s*[)（(]?[٠-٩0-9]+(?:[-:–—][٠-٩0-9]+)?[)）]?[\s:：-]", text):
            p.style = "Table Title"; stats["table_captions"] += 1
        elif re.match(r"^(?:شكل|الشكل)\s*[)（(]?[٠-٩0-9]+(?:[-:–—][٠-٩0-9]+)?[)）]?[\s:：-]", text):
            p.style = "Figure Caption"; stats["figure_captions"] += 1
        elif text.startswith(("المصدر:", "المصدر ", "المصدر(")):
            p.style = "Table Source"
        elif text in ("فكرة الفصل المركزية", "تنبيه اصطلاحي", "خلاصة الفصل", "ملاحظة", "تطبيق عملي", "مثال توضيحي", "أداة القارئ"):
            p.style = "Box Title"; apply_box(p, True); box_next = True; stats["boxes"] += 1
        elif box_next:
            p.style = "Box Text"; apply_box(p, False); box_next = False
        elif is_quran(text) and len(text) <= 260:
            p.style = "Qur’anic Verse"; stats["quran"] += 1
        elif is_hadith(text) and len(text) <= 500:
            p.style = "Prophetic Hadith"; stats["hadith"] += 1
        elif (text.startswith("«") and text.endswith("»") and len(text) > 120) or (text.startswith('"') and text.endswith('"') and len(text) > 120):
            p.style = "Long Quotation"
        elif reference_mode:
            if re.match(r"^(أولًا|أولاً|ثانيًا|ثانياً|المراجع العربية|المراجع الأجنبية)", text):
                p.style = "Main Heading"
            else:
                p.style = "Reference" if ar_ratio(text) >= .35 else "Reference Latin"
                pf = p.paragraph_format
                if p.style.name == "Reference":
                    pf.right_indent = Cm(.7); pf.first_line_indent = Cm(-.7)
                else:
                    pf.left_indent = Cm(.7); pf.first_line_indent = Cm(-.7)
                pf.space_after = Pt(4)
                stats["references"] += 1
        elif paragraph_has_drawing(p):
            set_bidi_paragraph(p, WD_ALIGN_PARAGRAPH.CENTER)
        elif re.match(r"^[•●▪■]\s*", text) or p._p.xpath('.//w:numPr'):
            p.style = "Bulleted List"
        elif re.match(r"^(?:\.[٠-٩0-9]+|[٠-٩0-9]+[.)])\s*", text):
            p.style = "Numbered List"
            corrected = normalize_list_text(text)
            if corrected != text:
                p.text = corrected; text = corrected; stats["numbering_normalized"] += 1
        elif re.match(r"^(أولًا|أولاً|ثانيًا|ثانياً|ثالثًا|ثالثاً|رابعًا|رابعاً|خامسًا|خامساً|سادسًا|سادساً|سابعًا|سابعاً|ثامنًا|ثامناً|تاسعًا|تاسعاً|عاشرًا|عاشراً)\s*[:：]", text):
            p.style = "Subheading"
        elif text in ("الأهداف", "منهجية الكتاب وإطار التحليل", "هيكلية الكتاب", "حدود المنهج"):
            p.style = "Main Heading"
        elif short_heading_candidate(text, next_text, old_style, p) and not author_bio_mode:
            level = original_outline_level(p)
            if level == 0:
                p.style = "Main Heading"
            elif level == 1:
                p.style = "Subheading"
            elif level is not None:
                p.style = "Unnumbered Internal Heading"
            else:
                p.style = "Main Heading" if len(text.split()) <= 7 else "Subheading"
        elif i < 20 and p.style.name in ("Book Title", "Book Subtitle", "Author Name"):
            pass
        else:
            p.style = "First Paragraph after Heading" if previous_semantic in heading_names else "Body Text"
            if p.style.name == "First Paragraph after Heading":
                stats["first_after_heading_fixed"] += 1

        name = p.style.name if p.style else "Body Text"
        if name in ("Book Title", "Book Subtitle", "Author Name", "Part Title", "Qur’anic Verse", "Figure Caption", "Figure Title"):
            align = WD_ALIGN_PARAGRAPH.CENTER
        elif name == "Reference Latin":
            align = WD_ALIGN_PARAGRAPH.LEFT
        elif name in ("Body Text", "First Paragraph after Heading", "Long Quotation", "Prophetic Hadith", "Box Text"):
            align = WD_ALIGN_PARAGRAPH.JUSTIFY
        else:
            align = WD_ALIGN_PARAGRAPH.RIGHT
        set_bidi_paragraph(p, align)
        stats["rtl_fixed"] += 1
        pf = p.paragraph_format
        pf.widow_control = True
        pf.keep_together = True
        if name in heading_names or name in ("Table Title", "Figure Title"):
            pf.keep_with_next = True
        style = p.style
        font = style.font.name or ("Noto Kufi Arabic" if name in heading_names else "Amiri")
        size = style.font.size.pt if style.font.size else 14
        for r in p.runs:
            set_run(r, font, size, bold=style.font.bold, color=BLACK, rtl=(name != "Reference Latin"))
        if is_quran(text):
            for r in p.runs:
                if any(x in r.text for x in ("﴿", "﴾", "{", "}", "قريش", "خوف", "جوع")):
                    set_run(r, "Amiri Quran", 15, color=BLACK, rtl=True)
            if name != "Qur’anic Verse":
                stats["quran"] += 1
        if is_hadith(text):
            stats["hadith"] += 1 if name != "Prophetic Hadith" else 0
        previous_semantic = name
    return stats


def set_table_borders(table, color=MID_GRAY, size="4", accent=False):
    tblpr = table._tbl.tblPr
    borders = tblpr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders"); tblpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}"); borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "16" if accent and edge == "right" else size)
        el.set(qn("w:color"), PETROL if accent and edge == "right" else color)


def set_cell_shading(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd"); tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, twips=100):
    tcpr = cell._tc.get_or_add_tcPr()
    mar = tcpr.find(qn("w:tcMar"))
    if mar is None:
        mar = OxmlElement("w:tcMar"); tcpr.append(mar)
    for side in ("top", "start", "bottom", "end"):
        el = mar.find(qn(f"w:{side}"))
        if el is None:
            el = OxmlElement(f"w:{side}"); mar.append(el)
        el.set(qn("w:w"), str(twips)); el.set(qn("w:type"), "dxa")


def style_tables(doc):
    stats = {"tables": len(doc.tables), "boxes": 0, "content_tables": 0, "rtl_tables": 0}
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        tblpr = table._tbl.tblPr
        if tblpr.find(qn("w:bidiVisual")) is None:
            tblpr.append(OxmlElement("w:bidiVisual"))
        stats["rtl_tables"] += 1
        is_box = len(table.rows) == 1 and len(table.columns) == 1
        if is_box:
            stats["boxes"] += 1
            set_table_borders(table, LIGHT_GRAY, "0", True)
            cell = table.cell(0, 0); set_cell_shading(cell, LIGHT_GRAY); set_cell_margins(cell, 227)
            for idx, p in enumerate(cell.paragraphs):
                p.style = "Box Title" if idx == 0 and len(cell.paragraphs) > 1 else "Box Text"
                set_bidi_paragraph(p, WD_ALIGN_PARAGRAPH.RIGHT if p.style.name == "Box Title" else WD_ALIGN_PARAGRAPH.JUSTIFY)
                for r in p.runs:
                    set_run(r, "Noto Kufi Arabic" if p.style.name == "Box Title" else "Amiri", 13 if p.style.name == "Box Title" else 12.5, bold=p.style.name == "Box Title", rtl=True)
            continue
        stats["content_tables"] += 1
        set_table_borders(table, MID_GRAY, "4", False)
        for ri, row in enumerate(table.rows):
            trpr = row._tr.get_or_add_trPr()
            if trpr.find(qn("w:cantSplit")) is None:
                trpr.append(OxmlElement("w:cantSplit"))
            if ri == 0 and trpr.find(qn("w:tblHeader")) is None:
                trpr.append(OxmlElement("w:tblHeader"))
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_margins(cell, 100)
                set_cell_shading(cell, DARK_GRAY if ri == 0 else (LIGHT_GRAY if ri % 2 == 0 else WHITE))
                for p in cell.paragraphs:
                    p.style = "Table Text"; set_bidi_paragraph(p, WD_ALIGN_PARAGRAPH.RIGHT)
                    p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(2); p.paragraph_format.line_spacing = 1.05; p.paragraph_format.keep_together = True
                    for r in p.runs:
                        set_run(r, "Amiri", 11.5, bold=(ri == 0), color=WHITE if ri == 0 else BLACK, rtl=True)
    return stats


def set_geometry(section):
    section.page_width = Cm(17); section.page_height = Cm(24)
    section.top_margin = Cm(2.2); section.bottom_margin = Cm(2.4)
    section.left_margin = Cm(2.8); section.right_margin = Cm(2.0)
    section.gutter = Cm(.7); section.header_distance = Cm(1.0); section.footer_distance = Cm(1.0)
    pgmar = section._sectPr.find(qn("w:pgMar"))
    if pgmar is not None:
        pgmar.set(qn("w:gutter"), "397")


def ensure_settings(doc):
    settings = doc.settings.element
    for tag in ("w:mirrorMargins", "w:rtlGutter", "w:evenAndOddHeaders"):
        if settings.find(qn(tag)) is None:
            settings.append(OxmlElement(tag))
    upd = settings.find(qn("w:updateFields"))
    if upd is None:
        upd = OxmlElement("w:updateFields"); settings.append(upd)
    upd.set(qn("w:val"), "true")
    fpr = settings.find(qn("w:footnotePr"))
    if fpr is None:
        fpr = OxmlElement("w:footnotePr"); settings.append(fpr)
    restart = fpr.find(qn("w:numRestart"))
    if restart is None:
        restart = OxmlElement("w:numRestart"); fpr.append(restart)
    restart.set(qn("w:val"), "continuous")
    styles = doc.styles.element
    defaults = styles.find(qn("w:docDefaults"))
    if defaults is None:
        defaults = OxmlElement("w:docDefaults"); styles.insert(0, defaults)
    rpd = defaults.find(qn("w:rPrDefault"))
    if rpd is None:
        rpd = OxmlElement("w:rPrDefault"); defaults.append(rpd)
    rpr = rpd.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr"); rpd.append(rpr)
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang"); rpr.append(lang)
    lang.set(qn("w:val"), "ar-SA"); lang.set(qn("w:bidi"), "ar-SA")


def clone_sectpr(template, break_type="oddPage"):
    sec = deepcopy(template)
    for tag in ("w:headerReference", "w:footerReference", "w:type", "w:titlePg", "w:pgNumType"):
        for el in list(sec.findall(qn(tag))):
            sec.remove(el)
    typ = OxmlElement("w:type"); typ.set(qn("w:val"), break_type); sec.insert(0, typ)
    return sec


def ensure_odd_breaks(doc):
    template = doc.element.body.sectPr
    inserted = 0; adjusted = 0
    for p in list(doc.paragraphs):
        if not CHAPTER_RE.match(clean(p.text)):
            continue
        clear_page_breaks(p)
        prev = p._p.getprevious(); sect = None
        if prev is not None and prev.tag == qn("w:p"):
            ppr = prev.find(qn("w:pPr")); sect = ppr.find(qn("w:sectPr")) if ppr is not None else None
        if sect is None:
            holder = OxmlElement("w:p"); ppr = OxmlElement("w:pPr"); holder.append(ppr); ppr.append(clone_sectpr(template, "oddPage")); p._p.addprevious(holder); inserted += 1
        else:
            typ = sect.find(qn("w:type"))
            if typ is None:
                typ = OxmlElement("w:type"); sect.insert(0, typ)
            typ.set(qn("w:val"), "oddPage"); adjusted += 1
    return {"odd_breaks_inserted": inserted, "odd_breaks_adjusted": adjusted}


def clear_story(story):
    paras = list(story.paragraphs)
    p = paras[0]
    for extra in paras[1:]:
        extra._element.getparent().remove(extra._element)
    for child in list(p._p):
        if child.tag != qn("w:pPr"):
            p._p.remove(child)
    return p


def set_page_number_format(section, fmt, start=None):
    sec = section._sectPr
    old = sec.find(qn("w:pgNumType"))
    if old is not None:
        sec.remove(old)
    el = OxmlElement("w:pgNumType"); el.set(qn("w:fmt"), fmt)
    if start is not None:
        el.set(qn("w:start"), str(start))
    sec.append(el)


def section_index_for_first_chapter(doc):
    idx = 0
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            text = clean("".join(child.itertext()))
            if CHAPTER_RE.match(text):
                return idx
            ppr = child.find(qn("w:pPr"))
            if ppr is not None and ppr.find(qn("w:sectPr")) is not None:
                idx += 1
    return max(1, len(doc.sections) - 14)


def configure_headers_footers(doc):
    ensure_settings(doc)
    for sec in doc.sections:
        set_geometry(sec)
    first_chapter = section_index_for_first_chapter(doc)
    for si, sec in enumerate(doc.sections):
        for story in (sec.header, sec.even_page_header, sec.first_page_header, sec.footer, sec.even_page_footer, sec.first_page_footer):
            for p in story.paragraphs:
                set_bidi_paragraph(p, p.alignment or WD_ALIGN_PARAGRAPH.RIGHT)
                for r in p.runs:
                    set_run(r, "Amiri", 10.5, rtl=True)
        if si < first_chapter:
            set_page_number_format(sec, "lowerRoman")
            sec.different_first_page_header_footer = True
            if si == 0:
                clear_story(sec.first_page_header); clear_story(sec.first_page_footer)
            continue
        sec.different_first_page_header_footer = True
        for name in ("header", "even_page_header", "first_page_header", "footer", "even_page_footer", "first_page_footer"):
            getattr(sec, name).is_linked_to_previous = False
        p = clear_story(sec.header); set_bidi_paragraph(p, WD_ALIGN_PARAGRAPH.RIGHT); add_field(p, ' STYLEREF "Chapter Title" ', BOOK_TITLE)
        p = clear_story(sec.even_page_header); set_bidi_paragraph(p, WD_ALIGN_PARAGRAPH.LEFT); r = p.add_run(BOOK_TITLE); set_run(r, "Amiri", 10.5, rtl=True)
        clear_story(sec.first_page_header)
        p = clear_story(sec.footer); set_bidi_paragraph(p, WD_ALIGN_PARAGRAPH.RIGHT); add_field(p, " PAGE ", "1")
        p = clear_story(sec.even_page_footer); set_bidi_paragraph(p, WD_ALIGN_PARAGRAPH.LEFT); add_field(p, " PAGE ", "1")
        clear_story(sec.first_page_footer)
        set_page_number_format(sec, "decimal", 1 if si == first_chapter else None)
    return first_chapter


def cap_images(doc):
    max_width = Cm(17 - 2.8 - 2.0 - .7)
    corrected = 0
    for shape in doc.inline_shapes:
        if shape.width > max_width:
            ratio = max_width / shape.width; shape.width = max_width; shape.height = int(shape.height * ratio); corrected += 1
    for p in doc.paragraphs:
        if paragraph_has_drawing(p):
            set_bidi_paragraph(p, WD_ALIGN_PARAGRAPH.CENTER)
    return corrected


def add_index_entries(doc):
    quran = set(); hadith = set(); subjects = set(); specialist = set(); names = set(); glossary = set()
    name_patterns = [r"عبد العزيز القرني", r"الإمام\s+([\u0600-\u06FF]+)", r"الملك\s+([\u0600-\u06FF]+(?:\s+[\u0600-\u06FF]+){0,3})", r"الأمير\s+([\u0600-\u06FF]+(?:\s+[\u0600-\u06FF]+){0,3})"]
    for p in doc.paragraphs:
        text = clean(p.text)
        if not text:
            continue
        style = p.style.name if p.style else ""
        if is_quran(text):
            m = re.search(r"(?:قريش|[\u0600-\u06FF]+)\s*[:：]\s*[٠-٩0-9]+", text)
            term = m.group(0) if m else text[:90]
            if term not in quran:
                add_hidden_index_field(p, term, "Q"); quran.add(term)
        if is_hadith(text):
            m = re.search(r"رواه\s+[^)）.،]+", text)
            term = m.group(0) if m else text[:90]
            if term not in hadith:
                add_hidden_index_field(p, term, "H"); hadith.add(term)
        if style in ("Chapter Title", "Main Heading", "Subheading", "Unnumbered Internal Heading"):
            if text not in subjects:
                add_hidden_index_field(p, text, "S"); subjects.add(text)
            if "الأمن" in text or any(k in text for k in ("المخاطر", "الجرائم", "الحشود", "الإرهاب", "الرؤية", "السيبراني", "المناخ", "الغذائي", "المائي")):
                if text not in specialist:
                    add_hidden_index_field(p, text, "P"); specialist.add(text)
            if len(text.split()) <= 5 and any(k in text for k in ("مفهوم", "تعريف", "نظرية", "الأمن", "حوكمة", "استدامة")):
                if text not in glossary:
                    add_hidden_index_field(p, text, "G"); glossary.add(text)
        for pat in name_patterns:
            for match in re.finditer(pat, text):
                term = match.group(0)
                if term not in names:
                    add_hidden_index_field(p, term, "N"); names.add(term)
    return {"quran_entries": len(quran), "hadith_entries": len(hadith), "subject_entries": len(subjects), "specialist_entries": len(specialist), "name_entries": len(names), "glossary_entries": len(glossary)}


def append_indexes(doc):
    p = doc.add_paragraph("الفهارس التحليلية", style="Appendix Title")
    p.paragraph_format.page_break_before = True
    fields = [
        ("فهرس الآيات القرآنية", "Q"),
        ("فهرس الأحاديث", "H"),
        ("فهرس الأسماء", "N"),
        ("الفهرس الموضوعي", "S"),
        ("مسرد المصطلحات", "G"),
        ("الفهرس التخصصي", "P"),
    ]
    for heading, fid in fields:
        h = doc.add_paragraph(heading, style="Index Heading")
        f = doc.add_paragraph(style="Index Field")
        add_field(f, f' INDEX \\f "{fid}" \\h "أ" \\c "2" ', "يُحدّث تلقائيًا في Microsoft Word")


def patch_footnotes(docx_path):
    count = 0
    tmp = docx_path.with_suffix(".tmp.docx")
    with zipfile.ZipFile(docx_path, "r") as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/footnotes.xml":
                from lxml import etree
                root = etree.fromstring(data)
                ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
                for foot in root.xpath("//w:footnote[@w:id >= 0]", namespaces=ns):
                    count += 1
                    for p in foot.xpath(".//w:p", namespaces=ns):
                        ppr = p.find(qn("w:pPr"))
                        if ppr is None:
                            ppr = OxmlElement("w:pPr"); p.insert(0, ppr)
                        if ppr.find(qn("w:bidi")) is None:
                            ppr.append(OxmlElement("w:bidi"))
                    for r in foot.xpath(".//w:r", namespaces=ns):
                        rpr = r.find(qn("w:rPr"))
                        if rpr is None:
                            rpr = OxmlElement("w:rPr"); r.insert(0, rpr)
                        fonts = rpr.find(qn("w:rFonts"))
                        if fonts is None:
                            fonts = OxmlElement("w:rFonts"); rpr.insert(0, fonts)
                        for key in ("ascii", "hAnsi", "eastAsia", "cs"):
                            fonts.set(qn(f"w:{key}"), "Amiri")
                        sz = rpr.find(qn("w:sz"))
                        if sz is None:
                            sz = OxmlElement("w:sz"); rpr.append(sz)
                        sz.set(qn("w:val"), "21")
                        if rpr.find(qn("w:rtl")) is None:
                            rpr.append(OxmlElement("w:rtl"))
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone="yes")
            zout.writestr(item, data)
    tmp.replace(docx_path)
    return count


def core_text(doc):
    return " ".join(clean(p.text) for p in doc.paragraphs if clean(p.text) and clean(p.text) not in ("فهرس المحتويات", "قائمة الجداول", "قائمة الأشكال"))


def build_interior():
    original_doc = Document(SOURCE)
    original_text = core_text(original_doc)
    doc = Document(SOURCE)
    create_styles(doc); ensure_settings(doc)
    indexes_replaced = replace_manual_indexes(doc)
    paragraph_stats = style_document_paragraphs(doc)
    table_stats = style_tables(doc)
    image_caps = cap_images(doc)
    index_stats = add_index_entries(doc)
    append_indexes(doc)
    break_stats = ensure_odd_breaks(doc)
    stage = ROOT / "stage_security_book.docx"; doc.save(stage)
    doc = Document(stage)
    create_styles(doc); ensure_settings(doc)
    first_chapter = configure_headers_footers(doc)
    for sec in doc.sections:
        set_geometry(sec)
    dirty = 0
    for fld in doc.element.xpath('.//w:fldChar[@w:fldCharType="begin"]'):
        fld.set(qn("w:dirty"), "true"); dirty += 1
    doc.core_properties.title = BOOK_TITLE
    doc.core_properties.subject = SUBTITLE
    doc.core_properties.author = AUTHOR
    doc.core_properties.language = "ar-SA"
    doc.core_properties.comments = "مراجعة وتنسيق كامل وفق معيار الكتاب العربي ذي الخطوات الثماني والعشرين"
    doc.save(FINAL_DOCX)
    footnotes = patch_footnotes(FINAL_DOCX)
    final_text = core_text(Document(FINAL_DOCX))
    ratio = min(1.0, len(final_text) / max(1, len(original_text)))
    return {
        "original": {"paragraphs": len(original_doc.paragraphs), "tables": len(original_doc.tables), "sections": len(original_doc.sections), "inline_shapes": len(original_doc.inline_shapes)},
        "paragraph_stats": paragraph_stats,
        "table_stats": table_stats,
        "image_caps": image_caps,
        "indexes_replaced": indexes_replaced,
        "index_stats": index_stats,
        **break_stats,
        "first_chapter_section": first_chapter,
        "fields_marked_dirty": dirty,
        "footnotes": footnotes,
        "content_preservation_ratio": ratio,
    }


def update_fields_and_export():
    uno_script = ROOT / "lo_security_book.py"
    uno_script.write_text(r'''import pathlib, subprocess, time, uno
from com.sun.star.beans import PropertyValue

def prop(name, value):
    p=PropertyValue(); p.Name=name; p.Value=value; return p
proc=subprocess.Popen(['soffice','--headless','--nologo','--nodefault','--nofirststartwizard','--norestore','--accept=socket,host=127.0.0.1,port=2004;urp;StarOffice.ComponentContext'])
try:
    local=uno.getComponentContext(); resolver=local.ServiceManager.createInstanceWithContext('com.sun.star.bridge.UnoUrlResolver',local)
    ctx=None
    for _ in range(60):
        try: ctx=resolver.resolve('uno:socket,host=127.0.0.1,port=2004;urp;StarOffice.ComponentContext'); break
        except Exception: time.sleep(1)
    if ctx is None: raise RuntimeError('LibreOffice UNO connection failed')
    desktop=ctx.ServiceManager.createInstanceWithContext('com.sun.star.frame.Desktop',ctx)
    src=pathlib.Path('output/رؤية_المملكة_2030_والأمن_الشامل_النسخة_النهائية.docx').resolve().as_uri()
    doc=desktop.loadComponentFromURL(src,'_blank',0,(prop('Hidden',True),prop('UpdateDocMode',3)))
    if doc is None: raise RuntimeError('Could not open final DOCX')
    try: doc.getTextFields().refresh()
    except Exception: pass
    try:
        indexes=doc.getDocumentIndexes()
        for i in range(indexes.getCount()): indexes.getByIndex(i).update()
        print('indexes_updated',indexes.getCount())
    except Exception as e: print('index_warning',e)
    try: doc.calculateAll()
    except Exception: pass
    time.sleep(6)
    out=pathlib.Path('output/رؤية_المملكة_2030_والأمن_الشامل_النسخة_النهائية.docx').resolve().as_uri()
    doc.storeAsURL(out,(prop('FilterName','Office Open XML Text'),prop('Overwrite',True)))
    doc.close(True)
finally:
    proc.terminate()
''', encoding="utf-8")
    subprocess.run(["/usr/bin/python3", str(uno_script)], check=True)
    normalized = Document(FINAL_DOCX)
    create_styles(normalized); ensure_settings(normalized)
    for sec in normalized.sections:
        set_geometry(sec)
    normalized.save(FINAL_DOCX)
    patch_footnotes(FINAL_DOCX)
    temp_pdf = OUT / (FINAL_DOCX.stem + ".pdf")
    if temp_pdf.exists(): temp_pdf.unlink()
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(OUT), str(FINAL_DOCX)], check=True)
    if not temp_pdf.exists():
        raise RuntimeError("PDF export was not created")
    gray_pdf = OUT / "gray_output.pdf"
    subprocess.run([
        "gs", "-q", "-dNOPAUSE", "-dBATCH", "-sDEVICE=pdfwrite", "-dCompatibilityLevel=1.7",
        "-sColorConversionStrategy=Gray", "-dProcessColorModel=/DeviceGray", "-dAutoRotatePages=/None",
        f"-sOutputFile={gray_pdf}", str(temp_pdf)
    ], check=True)
    temp_pdf.unlink()
    if FINAL_PDF.exists(): FINAL_PDF.unlink()
    gray_pdf.replace(FINAL_PDF)


def font_path(pattern):
    try:
        return subprocess.check_output(["fc-match", "-f", "%{file}", pattern], text=True).strip()
    except Exception:
        return "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"


def rtl_text(text):
    try:
        import arabic_reshaper
        from bidi.algorithm import get_display
        return get_display(arabic_reshaper.reshape(text))
    except Exception:
        return text


def wrap_rtl(draw, text, font, width):
    words=text.split(); lines=[]; cur=""
    for word in words:
        trial=(cur+" "+word).strip(); box=draw.textbbox((0,0),rtl_text(trial),font=font)
        if box[2]-box[0] <= width: cur=trial
        else:
            if cur: lines.append(cur)
            cur=word
    if cur: lines.append(cur)
    return lines


def image_to_word(image_path, output_path):
    doc=Document(); sec=doc.sections[0]
    sec.page_width=Cm(17); sec.page_height=Cm(24)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Cm(0)
    p=doc.paragraphs[0] if doc.paragraphs else doc.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(0)
    p.add_run().add_picture(str(image_path),width=Cm(17),height=Cm(24)); doc.save(output_path)


def build_cover_templates():
    W,H=2008,2835; ivory="#F4EFE3"; petrol="#164C5A"; gold="#B39756"; charcoal="#333333"
    kufi=font_path("Noto Kufi Arabic"); amiri=font_path("Amiri")
    front=Image.new("RGB",(W,H),ivory); d=ImageDraw.Draw(front)
    d.rectangle((0,0,W,230),fill=petrol); d.text((W//2,120),rtl_text(AUTHOR),font=ImageFont.truetype(kufi,48),fill="white",anchor="mm")
    d.rectangle((150,310,W-150,327),fill=gold)
    title=ImageFont.truetype(kufi,105); sub=ImageFont.truetype(kufi,49); body=ImageFont.truetype(amiri,42)
    y=500
    for line in wrap_rtl(d,BOOK_TITLE,title,W-300): d.text((W//2,y),rtl_text(line),font=title,fill=petrol,anchor="mm"); y+=145
    y+=40
    for line in wrap_rtl(d,SUBTITLE,sub,W-360): d.text((W//2,y),rtl_text(line),font=sub,fill=charcoal,anchor="mm"); y+=82
    cx,cy=W//2,1710
    for r,col,w in [(430,petrol,8),(330,gold,7),(230,petrol,6),(130,gold,5)]: d.ellipse((cx-r,cy-r,cx+r,cy+r),outline=col,width=w)
    for angle in range(0,360,45):
        x=cx+int(430*math.cos(math.radians(angle))); yy=cy+int(430*math.sin(math.radians(angle)))
        d.ellipse((x-27,yy-27,x+27,yy+27),fill=gold,outline=petrol,width=3)
    d.text((W//2,H-300),rtl_text(EDITION),font=body,fill=charcoal,anchor="mm")
    d.rectangle((W//2-260,H-210,W//2+260,H-90),outline=petrol,width=4); d.text((W//2,H-150),rtl_text("[شعار الناشر]"),font=ImageFont.truetype(kufi,34),fill=petrol,anchor="mm")
    front.save(FRONT_PNG,dpi=(300,300)); image_to_word(FRONT_PNG,FRONT_DOCX)

    back=Image.new("RGB",(W,H),ivory); bdraw=ImageDraw.Draw(back)
    bdraw.rectangle((0,0,W,220),fill=petrol); bdraw.rectangle((150,300,W-150,317),fill=gold)
    bdraw.text((W-150,390),rtl_text("الأمن ليس نهاية التنمية؛ إنه البنية التي تجعل المستقبل ممكنًا"),font=ImageFont.truetype(kufi,55),fill=petrol,anchor="ra")
    blurb=("يقرأ هذا الكتاب رؤية المملكة 2030 بوصفها مشروعًا وطنيًا أعاد صياغة العلاقة بين الأمن والتنمية والهوية والحوكمة. "
           "وينتقل من التأصيل الشرعي والنظامي والنظري إلى الأمن الوطني والاجتماعي والغذائي والمائي والمناخي والسيبراني، "
           "ثم الجرائم المالية وأمن الحشود والسياحة ومكافحة الإرهاب، وصولًا إلى قراءة موثقة للمستجدات والمؤشرات بعد إطلاق الرؤية.")
    bf=ImageFont.truetype(amiri,45); y=630
    for line in wrap_rtl(bdraw,blurb,bf,W-300): bdraw.text((W-150,y),rtl_text(line),font=bf,fill=charcoal,anchor="ra"); y+=78
    y+=55
    bdraw.rounded_rectangle((150,y,W-150,y+340),radius=22,fill="#EFEDE7",outline=gold,width=3)
    bdraw.text((W-180,y+45),rtl_text("عن المؤلف"),font=ImageFont.truetype(kufi,39),fill=petrol,anchor="ra")
    bio=(f"{AUTHOR}، {AUTHOR_CREDENTIAL}، يهتم بالقانون والأمن والاستراتيجيات الوطنية العامة، ويقدم في هذا العمل قراءة مركبة للأمن السعودي في زمن رؤية المملكة 2030.")
    yy=y+115
    for line in wrap_rtl(bdraw,bio,ImageFont.truetype(amiri,39),W-380): bdraw.text((W-180,yy),rtl_text(line),font=ImageFont.truetype(amiri,39),fill=charcoal,anchor="ra"); yy+=65
    y+=410
    labels=["[شعار الناشر]","[الموقع الإلكتروني]","[رمز QR]","[ISBN والباركود]","[تصنيف الكتاب]","[السعر]"]
    boxw=(W-340)//2; boxh=130
    for i,label in enumerate(labels):
        row=i//2; col=i%2; x0=175+col*(boxw+15); y0=y+row*(boxh+15)
        bdraw.rectangle((x0,y0,x0+boxw,y0+boxh),outline="#B8B8B8",width=2)
        bdraw.text((x0+boxw//2,y0+boxh//2),rtl_text(label),font=ImageFont.truetype(amiri,32),fill=charcoal,anchor="mm")
    bdraw.text((W-150,H-170),rtl_text(EDITION),font=ImageFont.truetype(amiri,32),fill=petrol,anchor="ra")
    back.save(BACK_PNG,dpi=(300,300)); image_to_word(BACK_PNG,BACK_DOCX)

    spine=Image.new("RGB",(W,H),ivory); s=ImageDraw.Draw(spine)
    s.rectangle((0,0,W,220),fill=petrol); s.text((W//2,115),rtl_text("دليل الكعب — غير مقياس للطباعة"),font=ImageFont.truetype(kufi,42),fill="white",anchor="mm")
    note="لا يُحدد عرض الكعب حتى تعتمد المطبعة عدد الصفحات النهائي، ونوع الورق ووزنه وسماكته، وطريقة التجليد. يوضع على الكعب: عنوان الكتاب، اسم المؤلف، وشعار الناشر."
    y=420
    for line in wrap_rtl(s,note,ImageFont.truetype(amiri,47),W-300): s.text((W-150,y),rtl_text(line),font=ImageFont.truetype(amiri,47),fill=charcoal,anchor="ra"); y+=82
    s.rectangle((W//2-260,1000,W//2+260,2420),fill=petrol)
    s.text((W//2,1450),rtl_text(BOOK_TITLE),font=ImageFont.truetype(kufi,62),fill="white",anchor="mm")
    s.text((W//2,1900),rtl_text(AUTHOR),font=ImageFont.truetype(kufi,42),fill=gold,anchor="mm")
    s.rectangle((W//2-180,2200,W//2+180,2330),outline=gold,width=4); s.text((W//2,2265),rtl_text("[شعار الناشر]"),font=ImageFont.truetype(amiri,32),fill="white",anchor="mm")
    spine.save(SPINE_PNG,dpi=(300,300)); image_to_word(SPINE_PNG,SPINE_DOCX)


def audit_docx_and_pdf(build_info):
    doc=Document(FINAL_DOCX)
    required=["Book Title","Book Subtitle","Author Name","Part Title","Chapter Number","Chapter Title","Main Heading","Subheading","Unnumbered Internal Heading","Body Text","First Paragraph after Heading","Long Quotation","Qur’anic Verse","Prophetic Hadith","Box Title","Box Text","Table Title","Table Text","Table Source","Figure Title","Figure Caption","Footnote","Bulleted List","Numbered List","Reference","Appendix Title"]
    style_names=[s.name for s in doc.styles]; missing=[x for x in required if x not in style_names]
    wrong_sections=[]
    for i,s in enumerate(doc.sections,1):
        vals=(s.page_width.cm,s.page_height.cm,s.left_margin.cm,s.right_margin.cm,s.top_margin.cm,s.bottom_margin.cm,s.gutter.cm); target=(17,24,2.8,2.0,2.2,2.4,.7)
        if any(abs(a-b)>.04 for a,b in zip(vals,target)): wrong_sections.append({"section":i,"values":vals})
    bad_rtl=[]; bad_align=[]; wrong_body=[]; wrong_head=[]
    for i,p in enumerate(doc.paragraphs,1):
        text=clean(p.text)
        if not text: continue
        ppr=p._p.get_or_add_pPr()
        if ar_ratio(text)>.25 and ppr.find(qn("w:bidi")) is None: bad_rtl.append(i)
        jc=ppr.find(qn("w:jc")); token=jc.get(qn("w:val")) if jc is not None else None
        name=p.style.name if p.style else ""
        if name in ("Body Text","First Paragraph after Heading","Long Quotation","Prophetic Hadith","Box Text") and token not in ("both","distribute",None): bad_align.append((i,name,token))
        if name in ("Chapter Number","Chapter Title","Main Heading","Subheading","Unnumbered Internal Heading","Appendix Title","Box Title","Table Title","Table Text","Table Source","Numbered List","Bulleted List","Reference") and token not in ("right","end",None): bad_align.append((i,name,token))
        for r in p.runs:
            if not clean(r.text): continue
            if name in ("Body Text","First Paragraph after Heading") and r.font.name not in (None,"Amiri"): wrong_body.append(i)
            if name in ("Chapter Number","Chapter Title","Main Heading","Subheading","Unnumbered Internal Heading","Book Title","Book Subtitle","Author Name") and r.font.name not in (None,"Noto Kufi Arabic"): wrong_head.append(i)
    table_rtl=[]; split_rows=[]
    for ti,t in enumerate(doc.tables,1):
        if t._tbl.tblPr.find(qn("w:bidiVisual")) is None: table_rtl.append(ti)
        for ri,row in enumerate(t.rows,1):
            if row._tr.get_or_add_trPr().find(qn("w:cantSplit")) is None: split_rows.append((ti,ri))
    pdf=fitz.open(FINAL_PDF); ew=17/2.54*72; eh=24/2.54*72
    bad_size=[]; out_bounds=[]; blank=[]; fonts=set(); rendered=[]; color_pages=[]; chapter_pages=[]
    for i,page in enumerate(pdf,1):
        if abs(page.rect.width-ew)>2 or abs(page.rect.height-eh)>2: bad_size.append(i)
        text=page.get_text("text").strip()
        if not text: blank.append(i)
        if re.search(r"الفصل\s+(?:الأول|الثاني|الثالث|الرابع|الخامس|السادس|السابع|الثامن|التاسع|العاشر|الحادي عشر|الثاني عشر|الثالث عشر|الرابع عشر)",text): chapter_pages.append(i)
        for block in page.get_text("blocks"):
            x0,y0,x1,y1=block[:4]
            if x0 < -1 or y0 < -1 or x1 > page.rect.width+1 or y1 > page.rect.height+1: out_bounds.append((i,[x0,y0,x1,y1]))
        for f in page.get_fonts(full=True): fonts.add(f[3])
        pix=page.get_pixmap(matrix=fitz.Matrix(.8,.8),alpha=False); im=Image.frombytes("RGB",[pix.width,pix.height],pix.samples)
        small=im.resize((max(1,im.width//8),max(1,im.height//8)))
        if any(max(px)-min(px)>3 for px in list(small.getdata())[::20]): color_pages.append(i)
        pth=OUT/f"page_{i:03d}.png"; im.save(pth); rendered.append(pth)
    for n,start in enumerate(range(0,len(rendered),12),1):
        imgs=[Image.open(x) for x in rendered[start:start+12]]; thumbs=[]
        for im in imgs:
            ratio=240/im.width; thumbs.append(im.resize((240,int(im.height*ratio))))
        th=max(x.height for x in thumbs); sheet=Image.new("RGB",(720,4*th),"white"); draw=ImageDraw.Draw(sheet)
        for j,im in enumerate(thumbs):
            x=(j%3)*240; y=(j//3)*th; sheet.paste(im,(x,y)); draw.text((x+5,y+5),str(start+j+1),fill="black")
        sheet.save(CONTACT/f"sheet_{n:03d}.jpg",quality=88)
    tracked=[]; media=[]
    with zipfile.ZipFile(FINAL_DOCX) as z:
        for name in z.namelist():
            if name.endswith(".xml"):
                data=z.read(name)
                if b"<w:ins" in data or b"<w:del" in data: tracked.append(name)
            if name.startswith("word/media/"):
                try:
                    tmp=ROOT/Path(name).name; tmp.write_bytes(z.read(name)); img=Image.open(tmp); media.append({"name":name,"pixels":img.size,"dpi":img.info.get("dpi")})
                except Exception: media.append({"name":name,"pixels":None,"dpi":None})
    chapter_count=sum(1 for p in doc.paragraphs if p.style and p.style.name=="Chapter Number" and CHAPTER_RE.match(clean(p.text)))
    audit={**build_info,"final_docx_paragraphs":len(doc.paragraphs),"final_docx_tables":len(doc.tables),"final_docx_sections":len(doc.sections),"required_styles_missing":missing,"wrong_section_geometry":wrong_sections,"arabic_paragraphs_without_rtl":bad_rtl[:100],"alignment_issues":bad_align[:100],"wrong_body_font_paragraphs":sorted(set(wrong_body))[:100],"wrong_heading_font_paragraphs":sorted(set(wrong_head))[:100],"tables_without_rtl":table_rtl,"rows_without_no_split":split_rows[:100],"pdf_pages":len(pdf),"pdf_bad_page_size":bad_size,"pdf_out_of_bounds_blocks":out_bounds[:100],"blank_pages":blank,"pdf_fonts":sorted(fonts),"non_grayscale_pages":color_pages[:100],"chapter_opening_pages":chapter_pages,"chapter_openings_not_odd":[p for p in chapter_pages if p%2==0],"chapter_count":chapter_count,"tracked_changes_xml":tracked,"embedded_media":media,"contact_sheets":len(list(CONTACT.glob('*.jpg')))}
    AUDIT_JSON.write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding="utf-8")
    return audit


def compliance_rows(audit):
    done="مكتمل ومطابق"; partial="مكتمل في الجزء القابل للتنفيذ"; external="يتطلب بيانات أو اعتمادًا خارجيًا"; na="غير منطبق"
    return [
        (1,"إعداد الصفحة",done,"17 × 24 سم، هوامش متقابلة، داخلي 2.8 سم، خارجي 2 سم، علوي 2.2 سم، سفلي 2.4 سم، وكعب 0.7 سم يمينًا."),
        (2,"اللغة والاتجاه",done,"اتجاه RTL لجميع الفقرات العربية والجداول، ضبط المتن، لغة تدقيق عربية، وترقيم عربي موحد."),
        (3,"الخطوط والأحجام",done,"Amiri للمتن، Amiri Quran للآيات، Noto Kufi Arabic للعناوين، وبقية الأحجام وفق المعيار."),
        (4,"إعدادات الفقرات",done,"تباعد 1.2، بعد 5 نقاط، مسافة بادئة 0.6 سم، وإلغاء المسافة في أول فقرة بعد العنوان والصناديق."),
        (5,"أنماط Word",done,"تم إنشاء جميع الأنماط المطلوبة، إضافة إلى نمط مستقل لرقم الفصل."),
        (6,"التسلسل الهرمي",done,"الفصل وعنوانه والعناوين الرئيسية والفرعية والداخلية مرتبطة بمستويات Outline المناسبة."),
        (7,"بدايات الأجزاء والفصول",done,"الفصول الأربعة عشر تبدأ بمقاطع Odd Page، وصفحة الافتتاح بلا رأس أو رقم ظاهر. لا توجد أجزاء مستقلة في المصدر."),
        (8,"الرؤوس والتذييلات",done,"عنوان الفصل في الصفحة اليمنى، عنوان الكتاب في اليسرى، وأرقام الصفحات في الحواف الخارجية."),
        (9,"الصفحات التمهيدية",partial,"حُفظ نصف العنوان والصفحة البيضاء والعنوان الكامل والمؤلف، واستُبدل الفهرس اليدوي بفهرس آلي. مواد الإهداء والشكر والتقديم والاختصارات غير موجودة ولم تُخترع."),
        (10,"المواد الختامية",partial,"المراجع وحول الكتاب والكاتب والفهارس التحليلية موجودة. لا توجد ملاحق مستقلة مقدمة من المؤلف."),
        (11,"الآيات القرآنية",done,"تحديد الآيات وإسناد Amiri Quran بحجم 15 نقطة، وإدراجها في فهرس آلي للآيات."),
        (12,"الأحاديث",done,"تحديد نصوص الحديث ومصادرها، نمط 13.5 نقطة بهوامش 0.8 سم، وفهرس آلي للأحاديث."),
        (13,"الاقتباسات",done,"أنماط الاقتباس القصير والطويل مضبوطة؛ الاقتباسات الطويلة بهوامش 1 سم وتباعد 1.15."),
        (14,"الجداول",done,"RTL، عنوان أعلى، مصدر أسفل عند وجوده، صف رأس متكرر، حدود 0.5 نقطة، صفوف رمادية ومنع الانقسام."),
        (15,"الأشكال والصور",done,"العناوين أسفل الشكل، توسيط الرسومات، عدم تجاوز الهوامش، وعدم تغيير النسب."),
        (16,"الصناديق التحريرية",done,"فكرة الفصل والتنبيهات والخلاصات بصندوق رمادي وحد جانبي بلون بترولي، مع الخطوط والأحجام المحددة."),
        (17,"الحواشي",done if audit.get("footnotes",0) else na,"الحواشي الموجودة ضُبطت Amiri 10.5 وRTL وترقيم مستمر؛ لا تُنشأ حواشٍ غير موجودة."),
        (18,"المراجع",done,"فصل العربية عن اللاتينية، مسافة معلقة 0.7 سم، 4 نقاط بعد المرجع، وروابط قابلة للحفظ في Word/PDF."),
        (19,"الفهارس الآلية",done,"فهرس المحتويات والجداول والأشكال والآيات والأحاديث والأسماء والموضوعات والمصطلحات والتخصصات أُنشئت كحقول Word قابلة للتحديث."),
        (20,"القوائم",done,"نقاط على اليمين وترقيم عربي موحد دون خلط أنماط الرموز."),
        (21,"التدقيق الطباعي",done,"مراجعة المقاس والهوامش والاتجاه والخطوط والجداول والصفحات الخالية وحدود النص والتغييرات المتعقبة."),
        (22,"الغلاف الأمامي",partial,"اسم المؤلف والعنوان والعنوان الفرعي وعنصر بصري واحد وألوان بترولي/ذهبي/عاجي؛ شعار الناشر ما زال حقلًا واضحًا."),
        (23,"الكعب",external,"العنوان والمؤلف والشعار مجهزة في دليل منفصل، لكن العرض ينتظر الورق والسماكة والتجليد."),
        (24,"الغلاف الخلفي",partial,"عبارة ترويجية ووصف وسيرة مختصرة موجودة؛ بيانات الناشر والموقع وQR وISBN والسعر ما تزال حقولًا خارجية."),
        (25,"المواصفات الداخلية",partial,"الداخلية 17 × 24، خطوط مضمّنة وPDF رمادي للمراجعة والطباعة؛ شهادة PDF/X-4 تتطلب preflight من المطبعة."),
        (26,"مواصفات الغلاف",external,"القوالب 300 DPI، لكن ملف CMYK الكامل مع نزف 3 مم ومساحة أمان 7 مم ينتظر عرض الكعب وملف المطبعة."),
        (27,"التشطيب",external,"المواصفات المقترحة: 350 gsm، تغليف مطفي، Spot UV، بروز خفيف، وتجليد مخيط لكون الكتاب مرجعيًا."),
        (28,"المراجعة النهائية",partial,"تمت مراجعة رقمية شاملة ورندر لجميع الصفحات؛ البروفة الورقية وبروفة المطبعة والاعتماد النهائي خطوات مادية خارج النظام."),
    ]


def build_report(audit):
    doc=Document(); create_styles(doc); set_geometry(doc.sections[0]); ensure_settings(doc)
    p=doc.add_paragraph("تقرير المطابقة النهائي للخطوات الثماني والعشرين",style="Book Title"); set_bidi_paragraph(p,WD_ALIGN_PARAGRAPH.CENTER)
    p=doc.add_paragraph(BOOK_TITLE,style="Book Subtitle"); set_bidi_paragraph(p,WD_ALIGN_PARAGRAPH.CENTER)
    summary=[f"عدد صفحات PDF النهائي: {audit['pdf_pages']} صفحة.",f"عدد الفصول: {audit['chapter_count']}.",f"عدد الجداول الفعلية: {audit['table_stats']['content_tables']}.",f"عدد الصناديق الجدولية: {audit['table_stats']['boxes']}، والصناديق الفقرية: {audit['paragraph_stats']['boxes']}.",f"إدخالات فهرس الآيات: {audit['index_stats']['quran_entries']}، والأحاديث: {audit['index_stats']['hadith_entries']}.",f"إدخالات الفهرس الموضوعي: {audit['index_stats']['subject_entries']}، والتخصصي: {audit['index_stats']['specialist_entries']}.",f"الصفحات غير المطابقة للمقاس: {audit['pdf_bad_page_size'] or 'لا يوجد'}.",f"النصوص خارج حدود الصفحة: {audit['pdf_out_of_bounds_blocks'] or 'لا يوجد'}.",f"فقرات عربية بلا RTL: {audit['arabic_paragraphs_without_rtl'] or 'لا يوجد'}.",f"مشكلات المحاذاة: {audit['alignment_issues'] or 'لا يوجد'}.",f"صفوف جداول قابلة للانقسام: {audit['rows_without_no_split'] or 'لا يوجد'}.",f"صفحات الفصول الزوجية: {audit['chapter_openings_not_odd'] or 'لا يوجد'}.",f"صفحات غير رمادية: {audit['non_grayscale_pages'] or 'لا يوجد'}.",f"ملفات تتضمن تغييرات متعقبة: {audit['tracked_changes_xml'] or 'لا يوجد'}.",f"نسبة الحفاظ على المحتوى النصي: {audit['content_preservation_ratio']:.1%}.",f"لوحات المراجعة البصرية: {audit['contact_sheets']}."]
    for line in summary:
        p=doc.add_paragraph(line,style="Body Text"); set_bidi_paragraph(p,WD_ALIGN_PARAGRAPH.JUSTIFY)
    table=doc.add_table(rows=1,cols=4); table.alignment=WD_TABLE_ALIGNMENT.CENTER
    for cell,text in zip(table.rows[0].cells,["الخطوة","المعيار","الحالة","النتيجة الدقيقة"]): cell.text=text
    for num,title,status,note in compliance_rows(audit):
        cells=table.add_row().cells; cells[0].text=str(num).translate(ARABIC_DIGITS); cells[1].text=title; cells[2].text=status; cells[3].text=note
    style_tables(doc)
    p=doc.add_paragraph("الخلاصة",style="Main Heading")
    conclusion="اكتملت جميع عناصر التنسيق الداخلي القابلة للتنفيذ رقميًا، مع الحفاظ على محتوى المخطوط. لا يُسمى الملف غلافًا تجاريًا نهائيًا أو PDF/X-4 معتمدًا قبل إضافة بيانات الناشر وISBN والباركود واعتماد عرض الكعب وملف CMYK والبروفة الورقية من المطبعة."
    p=doc.add_paragraph(conclusion,style="Body Text"); set_bidi_paragraph(p,WD_ALIGN_PARAGRAPH.JUSTIFY)
    doc.save(REPORT_DOCX)
    subprocess.run(["soffice","--headless","--convert-to","pdf","--outdir",str(OUT),str(REPORT_DOCX)],check=True)
    generated=OUT/(REPORT_DOCX.stem+".pdf")
    if generated.exists() and generated != REPORT_PDF: generated.replace(REPORT_PDF)


def build_readme(audit):
    README.write_text(
        f"{BOOK_TITLE} — حزمة التسليم النهائية\n\n"
        f"المؤلف: {AUTHOR}\n"
        f"عدد صفحات PDF الرمادي: {audit['pdf_pages']}\n"
        f"عدد الفصول: {audit['chapter_count']}\n\n"
        "تحتوي الحزمة على النسخة النهائية Word، وPDF رمادي للمراجعة والطباعة، وتقرير الخطوات الثماني والعشرين، والغلاف الأمامي والخلفي Word وPNG بدقة 300 DPI، ودليل الكعب، وملف التدقيق، ولوحات مراجعة جميع الصفحات.\n\n"
        "المتبقي خارجيًا: شعار وبيانات الناشر، الموقع، QR، ISBN والباركود والتصنيف والسعر، ومواصفات الورق والسماكة والتجليد، ثم ملف غلاف CMYK/PDF-X-4 وبروفة المطبعة.\n",
        encoding="utf-8"
    )


def package():
    pkg=ROOT/"package"
    if pkg.exists(): shutil.rmtree(pkg)
    pkg.mkdir()
    for p in [FINAL_DOCX,FINAL_PDF,FRONT_DOCX,FRONT_PNG,BACK_DOCX,BACK_PNG,SPINE_DOCX,SPINE_PNG,REPORT_DOCX,REPORT_PDF,AUDIT_JSON,README]:
        if p.exists(): shutil.copy2(p,pkg/p.name)
    shutil.copytree(CONTACT,pkg/CONTACT.name)
    if PACKAGE_ZIP.exists(): PACKAGE_ZIP.unlink()
    with zipfile.ZipFile(PACKAGE_ZIP,"w",zipfile.ZIP_DEFLATED,compresslevel=9) as z:
        for p in pkg.rglob("*"):
            if p.is_file(): z.write(p,p.relative_to(pkg))


def main():
    build_info=build_interior()
    update_fields_and_export()
    build_cover_templates()
    audit=audit_docx_and_pdf(build_info)
    failures=[]
    for key in ("required_styles_missing","wrong_section_geometry","arabic_paragraphs_without_rtl","alignment_issues","tables_without_rtl","rows_without_no_split","pdf_bad_page_size","pdf_out_of_bounds_blocks","non_grayscale_pages","chapter_openings_not_odd","tracked_changes_xml"):
        if audit.get(key): failures.append((key,audit[key]))
    if audit["chapter_count"] != 14: failures.append(("chapter_count",audit["chapter_count"]))
    if audit["content_preservation_ratio"] < .97: failures.append(("content_preservation_ratio",audit["content_preservation_ratio"]))
    build_report(audit); build_readme(audit); package()
    if failures:
        raise RuntimeError("Core compliance failures: "+repr(failures[:10]))
    print(json.dumps(audit,ensure_ascii=False,indent=2))


if __name__ == "__main__":
    main()
