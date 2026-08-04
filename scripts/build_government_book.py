#!/usr/bin/env python3
from __future__ import annotations

import json
import os
import re
import shutil
import subprocess
import textwrap
import zipfile
from copy import deepcopy
from pathlib import Path
from typing import Iterable, Optional

import fitz
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph

ROOT = Path.cwd()
SRC = ROOT / "source.docx"
OUT = ROOT / "output"
OUT.mkdir(exist_ok=True)
DOCX_OUT = OUT / "دوائر_الإنجاز_الحكومي_النسخة_النهائية.docx"
PDF_OUT = OUT / "دوائر_الإنجاز_الحكومي_النسخة_النهائية.pdf"
QA_DOCX = OUT / "تقرير_التدقيق_النهائي.docx"
QA_PDF = OUT / "تقرير_التدقيق_النهائي.pdf"
SPEC_DOCX = OUT / "مواصفات_الطباعة_والتجليد.docx"
SPEC_PDF = OUT / "مواصفات_الطباعة_والتجليد.pdf"
FRONT_COVER = OUT / "الغلاف_الأمامي_300DPI.png"
BACK_COVER = OUT / "الغلاف_الخلفي_300DPI.png"
SPINE_GUIDE = OUT / "دليل_الكعب_غير_مقاس.png"
CONTACT_DIR = OUT / "contact_sheets"
CONTACT_DIR.mkdir(exist_ok=True)
AUDIT_JSON = OUT / "audit.json"
README = OUT / "اقرأني_أولاً.txt"

BOOK_TITLE = "دوائر الإنجاز الحكومي"
SUBTITLE = "دليل عملي لحوكمة التحول وصناعة القرار والمساءلة عن النتائج في الجهات الحكومية السعودية"
TAGLINE = "من كثرة النشاط إلى أثر عام قابل للإثبات"
EDITION = "الطبعة الأولى — النسخة المنقحة والموسعة 2026م"
PETROL = "1D4E5F"
GOLD = "B39756"
LIGHT_GRAY = "F2F2F2"
MID_GRAY = "808080"
WHITE = "FFFFFF"
BLACK = "000000"

ARABIC_DIGITS = str.maketrans("0123456789", "٠١٢٣٤٥٦٧٨٩")


def ar_num(value: str | int) -> str:
    return str(value).translate(ARABIC_DIGITS)


def clean_text(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "").replace("\u200f", "").replace("\u200e", "")).strip()


def paragraph_text(p: Paragraph) -> str:
    return clean_text(p.text)


def has_arabic(s: str) -> bool:
    return bool(re.search(r"[\u0600-\u06FF]", s))


def arabic_ratio(s: str) -> float:
    letters = re.findall(r"[A-Za-z\u0600-\u06FF]", s)
    if not letters:
        return 0.0
    ar = sum(1 for x in letters if "\u0600" <= x <= "\u06FF")
    return ar / len(letters)


def set_run_font(run, font_name: str, size: float, bold: Optional[bool] = None, color: Optional[str] = None, rtl: bool = True):
    run.font.name = font_name
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._r.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attr}"), font_name)
    if rtl and rpr.find(qn("w:rtl")) is None:
        rpr.append(OxmlElement("w:rtl"))
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:val"), "ar-SA" if rtl else "en-US")
    lang.set(qn("w:bidi"), "ar-SA" if rtl else "en-US")


def set_para_rtl(p: Paragraph, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p.alignment = alignment
    ppr = p._p.get_or_add_pPr()
    if ppr.find(qn("w:bidi")) is None:
        ppr.insert(0, OxmlElement("w:bidi"))
    jc = ppr.find(qn("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc")
        ppr.append(jc)
    mapping = {
        WD_ALIGN_PARAGRAPH.RIGHT: "right",
        WD_ALIGN_PARAGRAPH.LEFT: "left",
        WD_ALIGN_PARAGRAPH.CENTER: "center",
        WD_ALIGN_PARAGRAPH.JUSTIFY: "both",
    }
    jc.set(qn("w:val"), mapping.get(alignment, "right"))


def set_keep(p: Paragraph, next_: bool = False, lines: bool = True, page_break: bool = False):
    pf = p.paragraph_format
    pf.keep_together = lines
    pf.keep_with_next = next_
    pf.page_break_before = page_break
    pf.widow_control = True


def style_font(style, font_name: str, size: float, bold: bool = False, color: str = BLACK):
    style.font.name = font_name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attr}"), font_name)
    if rpr.find(qn("w:rtl")) is None:
        rpr.append(OxmlElement("w:rtl"))
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:val"), "ar-SA")
    lang.set(qn("w:bidi"), "ar-SA")


def add_style(doc: Document, name: str, font: str, size: float, bold=False, align=WD_ALIGN_PARAGRAPH.RIGHT,
              before=0, after=5, line=1.2, first_indent=0, left_indent=0, right_indent=0,
              keep_next=False, keep_lines=True, outline: Optional[int] = None, color=BLACK):
    styles = doc.styles
    try:
        style = styles[name]
    except KeyError:
        style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style_font(style, font, size, bold, color)
    pf = style.paragraph_format
    pf.alignment = align
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.first_line_indent = Cm(first_indent) if first_indent else None
    pf.left_indent = Cm(left_indent) if left_indent else None
    pf.right_indent = Cm(right_indent) if right_indent else None
    pf.keep_with_next = keep_next
    pf.keep_together = keep_lines
    pf.widow_control = True
    ppr = style.element.get_or_add_pPr()
    if ppr.find(qn("w:bidi")) is None:
        ppr.insert(0, OxmlElement("w:bidi"))
    if outline is not None:
        out = ppr.find(qn("w:outlineLvl"))
        if out is None:
            out = OxmlElement("w:outlineLvl")
            ppr.append(out)
        out.set(qn("w:val"), str(outline))
    return style


def create_styles(doc: Document):
    add_style(doc, "Half Title", "Noto Kufi Arabic", 23, True, WD_ALIGN_PARAGRAPH.CENTER, before=100, after=0, line=1.0)
    add_style(doc, "Book Title", "Noto Kufi Arabic", 28, True, WD_ALIGN_PARAGRAPH.CENTER, before=70, after=20, line=1.05, color=PETROL)
    add_style(doc, "Book Subtitle", "Noto Kufi Arabic", 15, False, WD_ALIGN_PARAGRAPH.CENTER, before=4, after=5, line=1.2)
    add_style(doc, "Edition", "Amiri", 12.5, False, WD_ALIGN_PARAGRAPH.CENTER, before=10, after=0, line=1.2, color=MID_GRAY)
    add_style(doc, "Part Title", "Noto Kufi Arabic", 24, True, WD_ALIGN_PARAGRAPH.CENTER, before=95, after=16, line=1.05, keep_next=True, outline=0, color=PETROL)
    add_style(doc, "Part Subtitle", "Amiri", 16, False, WD_ALIGN_PARAGRAPH.CENTER, before=0, after=8, line=1.2, color=MID_GRAY)
    add_style(doc, "Chapter Title", "Noto Kufi Arabic", 22, True, WD_ALIGN_PARAGRAPH.RIGHT, before=80, after=22, line=1.1, keep_next=True, outline=1, color=PETROL)
    add_style(doc, "Main Heading", "Noto Kufi Arabic", 17, True, WD_ALIGN_PARAGRAPH.RIGHT, before=14, after=7, line=1.1, keep_next=True, outline=2, color=PETROL)
    add_style(doc, "Subheading", "Noto Kufi Arabic", 15, True, WD_ALIGN_PARAGRAPH.RIGHT, before=10, after=5, line=1.1, keep_next=True, outline=3, color=PETROL)
    add_style(doc, "Body Text Arabic", "Amiri", 14, False, WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=5, line=1.2, first_indent=.6)
    add_style(doc, "First Body Paragraph", "Amiri", 14, False, WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=5, line=1.2, first_indent=0)
    add_style(doc, "Numbered List Arabic", "Amiri", 13.5, False, WD_ALIGN_PARAGRAPH.RIGHT, before=0, after=4, line=1.18, left_indent=0, right_indent=.7)
    add_style(doc, "Table Title", "Amiri", 10.5, True, WD_ALIGN_PARAGRAPH.RIGHT, before=8, after=4, line=1.05, keep_next=True)
    add_style(doc, "Figure Caption", "Amiri", 10.5, False, WD_ALIGN_PARAGRAPH.CENTER, before=5, after=7, line=1.05, keep_next=True)
    add_style(doc, "Table Source", "Amiri", 10.5, False, WD_ALIGN_PARAGRAPH.RIGHT, before=3, after=7, line=1.0)
    add_style(doc, "Callout Title", "Noto Kufi Arabic", 13, True, WD_ALIGN_PARAGRAPH.RIGHT, before=7, after=3, line=1.1, keep_next=True, color=PETROL)
    add_style(doc, "Callout Text", "Amiri", 12.5, False, WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=8, line=1.2, first_indent=0)
    add_style(doc, "Reference Arabic", "Amiri", 11.5, False, WD_ALIGN_PARAGRAPH.RIGHT, before=0, after=4, line=1.15, left_indent=0, right_indent=.7)
    add_style(doc, "Reference Latin", "Amiri", 11.5, False, WD_ALIGN_PARAGRAPH.LEFT, before=0, after=4, line=1.15, left_indent=.7, right_indent=0)
    add_style(doc, "TOC Heading", "Noto Kufi Arabic", 20, True, WD_ALIGN_PARAGRAPH.RIGHT, before=25, after=12, line=1.05, color=PETROL)
    add_style(doc, "Front Heading", "Noto Kufi Arabic", 17, True, WD_ALIGN_PARAGRAPH.RIGHT, before=15, after=8, line=1.1, keep_next=True, color=PETROL)
    add_style(doc, "Publication Data", "Amiri", 12.5, False, WD_ALIGN_PARAGRAPH.RIGHT, before=0, after=4, line=1.15)


def add_field(p: Paragraph, instruction: str, placeholder="يُحدَّث تلقائيًا عند فتح الملف"):
    r1 = p.add_run()._r
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    r1.append(begin)
    r2 = p.add_run()._r
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    r2.append(instr)
    r3 = p.add_run()._r
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    r3.append(sep)
    p.add_run(placeholder)
    r4 = p.add_run()._r
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r4.append(end)


def insert_after(p: Paragraph, text: str = "", style: Optional[str] = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    p._p.addnext(new_p)
    new = Paragraph(new_p, p._parent)
    if text:
        new.add_run(text)
    if style:
        new.style = style
    return new


def delete_paragraph(p: Paragraph):
    parent = p._element.getparent()
    parent.remove(p._element)
    p._p = p._element = None


def replace_manual_indexes(doc: Document):
    paras = list(doc.paragraphs)
    toc_i = next((i for i, p in enumerate(paras) if clean_text(p.text) == "فهرس المحتويات"), None)
    part_i = next((i for i, p in enumerate(paras) if i > (toc_i or -1) and clean_text(p.text).startswith("الجزء الأول")), None)
    if toc_i is None or part_i is None or part_i <= toc_i:
        return False
    for p in paras[toc_i + 1:part_i]:
        delete_paragraph(p)
    toc = doc.paragraphs[toc_i]
    toc.style = "TOC Heading"
    field_p = insert_after(toc, style="Body Text Arabic")
    add_field(field_p, ' TOC \\o "1-3" \\h \\z \\u ')
    loh = insert_after(field_p, "قائمة الجداول", "TOC Heading")
    lot = insert_after(loh, style="Body Text Arabic")
    add_field(lot, ' TOC \\h \\z \\t "Table Title,1" ')
    lofh = insert_after(lot, "قائمة الأشكال", "TOC Heading")
    lof = insert_after(lofh, style="Body Text Arabic")
    add_field(lof, ' TOC \\h \\z \\t "Figure Caption,1" ')
    return True


def normalize_caption(text: str, kind: str) -> str:
    if kind == "table":
        m = re.match(r"^ال?جدول\s*[:：]?\s*([0-9٠-٩]+)\s*[—–-]?\s*(.*)$", text)
        if not m:
            m = re.match(r"^جدول\s+([0-9٠-٩]+)\s*[—–-]\s*(.*)$", text)
        if m:
            n = m.group(1).translate(ARABIC_DIGITS)
            rest = m.group(2).strip(" :—–-")
            return f"الجدول {n}: {rest}" if rest else f"الجدول {n}"
    else:
        m = re.match(r"^ال?شكل\s*[:：]?\s*([0-9٠-٩]+)\s*[—–-]?\s*(.*)$", text)
        if m:
            n = m.group(1).translate(ARABIC_DIGITS)
            rest = m.group(2).strip(" :—–-")
            return f"الشكل {n}: {rest}" if rest else f"الشكل {n}"
    return text


def heading_candidate(text: str, original_style: str, next_text: str) -> bool:
    if not text or len(text) > 105 or text.startswith("http") or text.startswith("["):
        return False
    if re.match(r"^[٠-٩0-9]+[.)-]", text):
        return False
    low = original_style.lower()
    if "heading" in low or "title" in low or "عنوان" in original_style:
        return True
    known = (
        "مشهد افتتاحي", "المشهد الأول", "قرار تحت المجهر", "حالة مركبة", "محادثة في غرفة القرار",
        "إشارة مبكرة", "اختبار واقعي", "من دفتر المحفظة", "لماذا دوائر لا مراحل؟", "قرار التوسع",
        "ورشة صياغة المشكلة", "قرار هذا الفصل", "الملحقات التطبيقية", "مسارات قراءة مقترحة",
        "وعد الكتاب", "لمن كتب؟", "خريطة الكتاب", "علاقة هذا الكتاب", "بيانات النشر", "تنبيه مهني",
        "دليل الانتقال إلى الأدوات المصاحبة", "المراجع والمصادر"
    )
    if any(text.startswith(x) for x in known):
        return True
    if text.endswith((".", "،", "؛")):
        return False
    words = text.split()
    if len(words) <= 11 and len(next_text) >= 90:
        common = ("في ", "قد ", "لا ", "يمكن ", "عند ", "تحتاج ", "يجب ", "هذا ", "هذه ", "إذا ", "كل ")
        return not text.startswith(common)
    return False


def apply_paragraph_shading(p: Paragraph, fill=LIGHT_GRAY, border_color=PETROL, border_size="18"):
    ppr = p._p.get_or_add_pPr()
    shd = ppr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        ppr.append(shd)
    shd.set(qn("w:fill"), fill)
    borders = ppr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        ppr.append(borders)
    right = borders.find(qn("w:right"))
    if right is None:
        right = OxmlElement("w:right")
        borders.append(right)
    right.set(qn("w:val"), "single")
    right.set(qn("w:sz"), border_size)
    right.set(qn("w:space"), "8")
    right.set(qn("w:color"), border_color)


def clear_page_breaks(p: Paragraph):
    ppr = p._p.get_or_add_pPr()
    pb = ppr.find(qn("w:pageBreakBefore"))
    if pb is not None:
        ppr.remove(pb)
    for br in list(p._p.xpath('.//w:br[@w:type="page"]')):
        br.getparent().remove(br)


def insert_section_before(p: Paragraph, template_sectpr, break_type="oddPage"):
    clear_page_breaks(p)
    blank = OxmlElement("w:p")
    p._p.addprevious(blank)
    ppr = OxmlElement("w:pPr")
    blank.append(ppr)
    sectpr = deepcopy(template_sectpr)
    for tag in ("w:headerReference", "w:footerReference", "w:pgNumType", "w:titlePg", "w:type"):
        for el in list(sectpr.findall(qn(tag))):
            sectpr.remove(el)
    typ = OxmlElement("w:type")
    typ.set(qn("w:val"), break_type)
    sectpr.insert(0, typ)
    ppr.append(sectpr)


def style_paragraphs(doc: Document):
    paras = list(doc.paragraphs)
    title_indices = [i for i, p in enumerate(paras) if paragraph_text(p) == BOOK_TITLE]
    first_title = title_indices[0] if title_indices else None
    full_title = title_indices[1] if len(title_indices) > 1 else first_title
    ref_mode = False
    after_heading = False
    callout_next = False
    part_subtitle_remaining = 0
    stats = {"parts": 0, "chapters": 0, "main_headings": 0, "tables_captions": 0, "figure_captions": 0, "references": 0, "callouts": 0}

    nonempty = [paragraph_text(p) for p in paras]
    for i, p in enumerate(paras):
        t = paragraph_text(p)
        if not t:
            p.paragraph_format.space_after = Pt(0)
            continue
        nxt = next((nonempty[j] for j in range(i + 1, len(nonempty)) if nonempty[j]), "")
        orig = p.style.name if p.style else ""
        set_para_rtl(p)

        if i == first_title:
            p.style = "Half Title"
            set_para_rtl(p, WD_ALIGN_PARAGRAPH.CENTER)
            after_heading = True
            continue
        if i == full_title and i != first_title:
            p.style = "Book Title"
            set_para_rtl(p, WD_ALIGN_PARAGRAPH.CENTER)
            after_heading = True
            continue
        if full_title is not None and full_title < i < full_title + 5 and t not in ("بيانات النشر",):
            p.style = "Edition" if "2026" in t or "الطبعة" in t else "Book Subtitle"
            set_para_rtl(p, WD_ALIGN_PARAGRAPH.CENTER)
            continue

        if t == "فهرس المحتويات" or t in ("قائمة الجداول", "قائمة الأشكال"):
            p.style = "TOC Heading"
            set_para_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
            after_heading = True
            continue

        if re.match(r"^الجزء\s+(الأول|الثاني|الثالث|الرابع|الخامس|السادس|السابع|الثامن|التاسع|العاشر|[0-9٠-٩]+)", t):
            p.style = "Part Title"
            set_para_rtl(p, WD_ALIGN_PARAGRAPH.CENTER)
            stats["parts"] += 1
            part_subtitle_remaining = 2
            after_heading = True
            continue
        if part_subtitle_remaining > 0:
            p.style = "Part Subtitle"
            set_para_rtl(p, WD_ALIGN_PARAGRAPH.CENTER)
            part_subtitle_remaining -= 1
            continue

        if re.match(r"^الفصل\s+[0-9٠-٩]+\s*[:：]", t) or t.startswith("خاتمة:") or re.match(r"^الملحق\s+", t) or t in ("الملحقات التطبيقية", "دليل الانتقال إلى الأدوات المصاحبة", "المراجع والمصادر"):
            p.style = "Chapter Title"
            set_para_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
            if t.startswith("الفصل"):
                stats["chapters"] += 1
            ref_mode = t == "المراجع والمصادر"
            after_heading = True
            continue

        if re.match(r"^جدول\s+[0-9٠-٩]+\s*[—–-]", t) or re.match(r"^الجدول\s*[:：]?\s*[0-9٠-٩]+", t):
            p.text = normalize_caption(t, "table")
            p.style = "Table Title"
            set_para_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
            stats["tables_captions"] += 1
            after_heading = True
            continue
        if re.match(r"^شكل\s+[0-9٠-٩]+\s*[—–-]", t) or re.match(r"^الشكل\s*[:：]?\s*[0-9٠-٩]+", t):
            p.text = normalize_caption(t, "figure")
            p.style = "Figure Caption"
            set_para_rtl(p, WD_ALIGN_PARAGRAPH.CENTER)
            stats["figure_captions"] += 1
            after_heading = False
            continue
        if t.startswith("المصدر:"):
            p.style = "Table Source"
            set_para_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
            continue

        if t == "قرار هذا الفصل":
            p.style = "Callout Title"
            set_para_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
            apply_paragraph_shading(p)
            callout_next = True
            stats["callouts"] += 1
            continue
        if t.startswith("الأدوات المرتبطة في الدليل المصاحب"):
            p.style = "Callout Text"
            set_para_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
            apply_paragraph_shading(p)
            stats["callouts"] += 1
            continue
        if callout_next:
            p.style = "Callout Text"
            set_para_rtl(p, WD_ALIGN_PARAGRAPH.JUSTIFY)
            apply_paragraph_shading(p)
            callout_next = False
            continue

        if t in ("بيانات النشر", "تنبيه مهني"):
            p.style = "Front Heading"
            set_para_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
            after_heading = True
            continue

        if ref_mode:
            if t.startswith("http://") or t.startswith("https://"):
                p.style = "Reference Latin"
                set_para_rtl(p, WD_ALIGN_PARAGRAPH.LEFT)
            elif re.match(r"^\[[0-9٠-٩]+\]", t):
                p.style = "Reference Arabic" if arabic_ratio(t) >= .35 else "Reference Latin"
                set_para_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT if arabic_ratio(t) >= .35 else WD_ALIGN_PARAGRAPH.LEFT)
                stats["references"] += 1
            elif heading_candidate(t, orig, nxt):
                p.style = "Main Heading"
                set_para_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
            else:
                p.style = "Reference Arabic" if arabic_ratio(t) >= .35 else "Reference Latin"
                set_para_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT if arabic_ratio(t) >= .35 else WD_ALIGN_PARAGRAPH.LEFT)
            continue

        if heading_candidate(t, orig, nxt):
            p.style = "Main Heading"
            set_para_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
            stats["main_headings"] += 1
            after_heading = True
            continue

        if re.match(r"^[٠-٩0-9]+[.)]\s*", t) or orig.lower().startswith("list"):
            p.style = "Numbered List Arabic"
            set_para_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
            after_heading = False
            continue

        p.style = "First Body Paragraph" if after_heading else "Body Text Arabic"
        set_para_rtl(p, WD_ALIGN_PARAGRAPH.JUSTIFY)
        after_heading = False

    for p in doc.paragraphs:
        st = p.style.name if p.style else ""
        if st in ("Part Title", "Chapter Title", "Main Heading", "Subheading", "Front Heading", "TOC Heading", "Table Title", "Callout Title"):
            set_keep(p, next_=True, lines=True)
        else:
            set_keep(p, next_=False, lines=True)
        for r in p.runs:
            rtl = arabic_ratio(r.text) >= .25 or st not in ("Reference Latin",)
            if st in ("Book Title", "Half Title", "Part Title", "Chapter Title", "Main Heading", "Subheading", "Front Heading", "TOC Heading", "Callout Title"):
                size = float(p.style.font.size.pt) if p.style.font.size else 14
                set_run_font(r, "Noto Kufi Arabic", size, bold=p.style.font.bold, color=PETROL, rtl=True)
            elif st == "Reference Latin":
                set_run_font(r, "Amiri", 11.5, rtl=False)
            else:
                size = float(p.style.font.size.pt) if p.style.font.size else 14
                set_run_font(r, "Amiri", size, bold=p.style.font.bold, rtl=rtl)
    return stats


def set_cell_shading(cell, fill: str):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90):
    tc = cell._tc
    tcpr = tc.get_or_add_tcPr()
    mar = tcpr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tcpr.append(mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        el = mar.find(qn(f"w:{m}"))
        if el is None:
            el = OxmlElement(f"w:{m}")
            mar.append(el)
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")


def style_tables(doc: Document):
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        tblpr = table._tbl.tblPr
        if tblpr.find(qn("w:bidiVisual")) is None:
            tblpr.append(OxmlElement("w:bidiVisual"))
        borders = tblpr.find(qn("w:tblBorders"))
        if borders is None:
            borders = OxmlElement("w:tblBorders")
            tblpr.append(borders)
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = borders.find(qn(f"w:{edge}"))
            if e is None:
                e = OxmlElement(f"w:{edge}")
                borders.append(e)
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), "4")
            e.set(qn("w:color"), MID_GRAY)
        for r_i, row in enumerate(table.rows):
            trpr = row._tr.get_or_add_trPr()
            if trpr.find(qn("w:cantSplit")) is None:
                trpr.append(OxmlElement("w:cantSplit"))
            if r_i == 0 and trpr.find(qn("w:tblHeader")) is None:
                trpr.append(OxmlElement("w:tblHeader"))
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_margins(cell)
                set_cell_shading(cell, PETROL if r_i == 0 else (LIGHT_GRAY if r_i % 2 == 0 else WHITE))
                for p in cell.paragraphs:
                    set_para_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.line_spacing = 1.05
                    for run in p.runs:
                        set_run_font(run, "Amiri", 11.5, bold=(r_i == 0), color=WHITE if r_i == 0 else BLACK, rtl=True)


def add_hyperlink(paragraph: Paragraph, url: str, display: Optional[str] = None):
    display = display or url
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), PETROL)
    rpr.append(color)
    run.append(rpr)
    txt = OxmlElement("w:t")
    txt.text = display
    run.append(txt)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def make_urls_clickable(doc: Document):
    count = 0
    for p in doc.paragraphs:
        t = paragraph_text(p)
        if t.startswith("http://") or t.startswith("https://"):
            for child in list(p._p):
                if child.tag != qn("w:pPr"):
                    p._p.remove(child)
            add_hyperlink(p, t, t)
            count += 1
    return count


def set_page_geometry(section):
    section.page_width = Cm(17)
    section.page_height = Cm(24)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.4)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.8)
    section.gutter = Cm(0.7)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)
    sectpr = section._sectPr
    pgmar = sectpr.find(qn("w:pgMar"))
    if pgmar is not None:
        pgmar.set(qn("w:gutter"), str(int(Cm(.7))))
    if sectpr.find(qn("w:mirrorMargins")) is None:
        sectpr.append(OxmlElement("w:mirrorMargins"))


def set_page_numbering(section, fmt: str, start: Optional[int] = None):
    sectpr = section._sectPr
    old = sectpr.find(qn("w:pgNumType"))
    if old is not None:
        sectpr.remove(old)
    el = OxmlElement("w:pgNumType")
    el.set(qn("w:fmt"), fmt)
    if start is not None:
        el.set(qn("w:start"), str(start))
    sectpr.append(el)


def clear_story(story):
    for p in list(story.paragraphs)[1:]:
        p._element.getparent().remove(p._element)
    p = story.paragraphs[0]
    for child in list(p._p):
        if child.tag != qn("w:pPr"):
            p._p.remove(child)
    return p


def add_page_field(p: Paragraph):
    add_field(p, " PAGE ", "1")


def add_styleref_field(p: Paragraph, style_name="Chapter Title"):
    add_field(p, f' STYLEREF "{style_name}" ', BOOK_TITLE)


def configure_sections(doc: Document):
    settings = doc.settings.element
    for name in ("mirrorMargins", "evenAndOddHeaders", "updateFields"):
        if settings.find(qn(f"w:{name}")) is None:
            el = OxmlElement(f"w:{name}")
            if name == "updateFields":
                el.set(qn("w:val"), "true")
            settings.append(el)
    sections = list(doc.sections)
    for s in sections:
        set_page_geometry(s)

    if sections:
        s0 = sections[0]
        s0.different_first_page_header_footer = True
        for story in (s0.header, s0.even_page_header, s0.first_page_header, s0.footer, s0.even_page_footer, s0.first_page_footer):
            clear_story(story)

    if len(sections) > 1:
        sf = sections[1]
        sf.different_first_page_header_footer = False
        sf.header.is_linked_to_previous = False
        sf.even_page_header.is_linked_to_previous = False
        sf.footer.is_linked_to_previous = False
        sf.even_page_footer.is_linked_to_previous = False
        clear_story(sf.header)
        clear_story(sf.even_page_header)
        p = clear_story(sf.footer)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        add_page_field(p)
        p = clear_story(sf.even_page_footer)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_page_field(p)
        set_page_numbering(sf, "lowerRoman", 5)

    for idx, s in enumerate(sections[2:], start=2):
        s.different_first_page_header_footer = True
        s.header.is_linked_to_previous = False
        s.even_page_header.is_linked_to_previous = False
        s.first_page_header.is_linked_to_previous = False
        s.footer.is_linked_to_previous = False
        s.even_page_footer.is_linked_to_previous = False
        s.first_page_footer.is_linked_to_previous = False
        p = clear_story(s.header)
        set_para_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
        add_styleref_field(p)
        for r in p.runs:
            set_run_font(r, "Amiri", 10.5, rtl=True)
        p = clear_story(s.even_page_header)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.add_run(BOOK_TITLE)
        for r in p.runs:
            set_run_font(r, "Amiri", 10.5, rtl=True)
        clear_story(s.first_page_header)
        p = clear_story(s.footer)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        add_page_field(p)
        p = clear_story(s.even_page_footer)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_page_field(p)
        clear_story(s.first_page_footer)
        set_page_numbering(s, "decimal", 1 if idx == 2 else None)


def insert_required_section_breaks(doc: Document):
    body_sectpr = doc.element.body.sectPr
    targets = []
    for p in doc.paragraphs:
        t = paragraph_text(p)
        if t == "تنبيه مهني":
            targets.append((p, "nextPage"))
        elif re.match(r"^الجزء\s+", t):
            targets.append((p, "oddPage"))
        elif re.match(r"^الفصل\s+[0-9٠-٩]+\s*[:：]", t) or t.startswith("خاتمة:") or re.match(r"^الملحق\s+", t) or t in ("الملحقات التطبيقية", "دليل الانتقال إلى الأدوات المصاحبة", "المراجع والمصادر"):
            targets.append((p, "oddPage"))
    seen = set()
    for p, kind in targets:
        key = id(p._p)
        if key in seen:
            continue
        insert_section_before(p, body_sectpr, kind)
        seen.add(key)
    return len(targets)


def set_core_properties(doc: Document):
    cp = doc.core_properties
    cp.title = BOOK_TITLE
    cp.subject = SUBTITLE
    cp.comments = "نسخة منسقة للطباعة — اتجاه عربي من اليمين إلى اليسار"
    cp.language = "ar-SA"


def build_book():
    if not SRC.exists():
        raise FileNotFoundError(SRC)
    original = Document(SRC)
    original_text = "\n".join(clean_text(p.text) for p in original.paragraphs if clean_text(p.text))
    original_counts = {
        "paragraphs": len(original.paragraphs),
        "tables": len(original.tables),
        "inline_shapes": len(original.inline_shapes),
        "characters": len(original_text),
    }
    doc = Document(SRC)
    create_styles(doc)
    replaced = replace_manual_indexes(doc)
    tmp1 = ROOT / "stage1.docx"
    doc.save(tmp1)
    doc = Document(tmp1)
    create_styles(doc)
    stats = style_paragraphs(doc)
    style_tables(doc)
    hyperlinks = make_urls_clickable(doc)
    set_core_properties(doc)
    breaks = insert_required_section_breaks(doc)
    tmp2 = ROOT / "stage2.docx"
    doc.save(tmp2)
    doc = Document(tmp2)
    create_styles(doc)
    configure_sections(doc)
    for p in doc.paragraphs:
        t = paragraph_text(p)
        if t in ("فهرس المحتويات", "قائمة الجداول", "قائمة الأشكال"):
            p.style = "TOC Heading"
    doc.save(DOCX_OUT)
    return {
        "original": original_counts,
        "styled": stats,
        "manual_index_replaced": replaced,
        "hyperlinks_created": hyperlinks,
        "section_breaks_inserted": breaks,
        "final_sections": len(Document(DOCX_OUT).sections),
        "final_tables": len(Document(DOCX_OUT).tables),
    }


def font_file(pattern: str) -> str:
    try:
        return subprocess.check_output(["fc-match", "-f", "%{file}", pattern], text=True).strip()
    except Exception:
        return "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"


def rtl_text(text: str) -> str:
    try:
        import arabic_reshaper
        from bidi.algorithm import get_display
        return get_display(arabic_reshaper.reshape(text))
    except Exception:
        return text


def wrap_ar(draw, text, font, max_width):
    words = text.split()
    lines, cur = [], ""
    for word in words:
        trial = (cur + " " + word).strip()
        bbox = draw.textbbox((0, 0), rtl_text(trial), font=font)
        if bbox[2] - bbox[0] <= max_width:
            cur = trial
        else:
            if cur:
                lines.append(cur)
            cur = word
    if cur:
        lines.append(cur)
    return lines


def make_covers():
    W, H = 2008, 2835
    ivory = "#F4EFE3"
    petrol = "#1D4E5F"
    gold = "#B39756"
    charcoal = "#333333"
    kufi = font_file("Noto Kufi Arabic")
    amiri = font_file("Amiri")
    title_f = ImageFont.truetype(kufi, 112)
    sub_f = ImageFont.truetype(kufi, 43)
    body_f = ImageFont.truetype(amiri, 47)
    small_f = ImageFont.truetype(amiri, 34)

    img = Image.new("RGB", (W, H), ivory)
    d = ImageDraw.Draw(img)
    d.rectangle((0, 0, W, 230), fill=petrol)
    d.rectangle((150, 310, W - 150, 327), fill=gold)
    cx, cy = W // 2, 1250
    radii = [410, 330, 250, 170]
    for i, r in enumerate(radii):
        d.ellipse((cx-r, cy-r, cx+r, cy+r), outline=petrol if i % 2 == 0 else gold, width=8)
    for angle in range(0, 360, 45):
        import math
        x = cx + int(410 * math.cos(math.radians(angle)))
        y = cy + int(410 * math.sin(math.radians(angle)))
        d.ellipse((x-32, y-32, x+32, y+32), fill=gold, outline=petrol, width=4)
    d.text((W-150, 405), rtl_text(BOOK_TITLE), font=title_f, fill=petrol, anchor="ra")
    y = 590
    for line in wrap_ar(d, SUBTITLE, sub_f, W-300):
        d.text((W-150, y), rtl_text(line), font=sub_f, fill=charcoal, anchor="ra")
        y += 72
    d.text((W-150, 770), rtl_text(TAGLINE), font=body_f, fill=gold, anchor="ra")
    d.text((W-150, H-270), rtl_text(EDITION), font=small_f, fill=petrol, anchor="ra")
    img.save(FRONT_COVER, dpi=(300, 300))

    back = Image.new("RGB", (W, H), ivory)
    b = ImageDraw.Draw(back)
    b.rectangle((0, 0, W, 220), fill=petrol)
    b.rectangle((150, 300, W-150, 317), fill=gold)
    b.text((W-150, 390), rtl_text("من كثرة النشاط إلى أثر عام قابل للإثبات"), font=ImageFont.truetype(kufi, 62), fill=petrol, anchor="ra")
    blurb = ("قد تمتلئ الجهة الحكومية بالمبادرات والاجتماعات ولوحات المتابعة، بينما تبقى النتيجة العامة بعيدة. "
             "يقدم هذا الكتاب نظامًا عمليًا يربط التفويض والشرعية بالنتيجة العامة، ويحوّل الأولويات إلى قرارات وبوابات، "
             "ويحدد حقوق القرار والاعتماديات، ثم يتحقق من التبنّي والمنافع ويعيد التعلم إلى المحفظة. صُمم للقيادات، "
             "وملاك المحافظ والبرامج، ومكاتب الاستراتيجية والإنجاز والتحول، والمالية والبيانات والمخاطر والمراجعة.")
    y = 640
    for line in wrap_ar(b, blurb, body_f, W-300):
        b.text((W-150, y), rtl_text(line), font=body_f, fill=charcoal, anchor="ra")
        y += 82
    y += 70
    bullets = [
        "تشخيص وهم النشاط وربط المبادرات بالنتائج العامة",
        "بناء حقوق القرار والمحفظة والبوابات المتناسبة",
        "تشغيل مكتب إنجاز يزيل العوائق بدل جمع التقارير",
        "إثبات المنافع وحدود الادعاء وإعادة التعلم إلى القرار",
    ]
    for item in bullets:
        b.ellipse((W-182, y+20, W-158, y+44), fill=gold)
        b.text((W-215, y), rtl_text(item), font=body_f, fill=charcoal, anchor="ra")
        y += 112
    b.rectangle((150, H-350, W-150, H-333), fill=gold)
    b.text((W-150, H-275), rtl_text(EDITION), font=small_f, fill=petrol, anchor="ra")
    back.save(BACK_COVER, dpi=(300, 300))

    sw = 1000
    spine = Image.new("RGB", (sw, H), ivory)
    s = ImageDraw.Draw(spine)
    s.rectangle((0, 0, sw, 220), fill=petrol)
    s.text((sw-80, 330), rtl_text("دليل إعداد الكعب"), font=ImageFont.truetype(kufi, 70), fill=petrol, anchor="ra")
    s.text((sw-80, 480), rtl_text("غير مقياس للطباعة"), font=ImageFont.truetype(kufi, 38), fill=gold, anchor="ra")
    note = "يُحسب عرض الكعب النهائي بعد اعتماد عدد صفحات ملف الطباعة، ونوع الورق وسماكته، وطريقة التجليد لدى المطبعة."
    y = 680
    for line in wrap_ar(s, note, ImageFont.truetype(amiri, 43), sw-160):
        s.text((sw-80, y), rtl_text(line), font=ImageFont.truetype(amiri, 43), fill=charcoal, anchor="ra")
        y += 75
    s.rectangle((80, 1200, sw-80, 1215), fill=gold)
    s.text((sw//2, 1500), rtl_text(BOOK_TITLE), font=ImageFont.truetype(kufi, 68), fill=petrol, anchor="mm")
    spine.save(SPINE_GUIDE, dpi=(300, 300))


def make_support_docs():
    spec = Document()
    create_styles(spec)
    sec = spec.sections[0]
    set_page_geometry(sec)
    p = spec.add_paragraph("مواصفات الطباعة والتجليد", style="Book Title")
    set_para_rtl(p, WD_ALIGN_PARAGRAPH.CENTER)
    items = [
        "مقاس القطع النهائي: 17 × 24 سم.",
        "الصفحات متقابلة، والهامش الداخلي 2.8 سم، والخارجي 2 سم، والعلوي 2.2 سم، والسفلي 2.4 سم، وهامش التجليد 0.7 سم.",
        "المتن: Amiri بحجم 14 نقطة، وتباعد 1.2، ومحاذاة كاملة من اليمين إلى اليسار.",
        "العناوين: Noto Kufi Arabic بأحجام هرمية ثابتة.",
        "الغلاف: 300 DPI، مع نزف 3 مم ومنطقة أمان لا تقل عن 7 مم عند إعداد الملف النهائي لدى المطبعة.",
        "التشطيب المقترح: غلاف 350 gsm، تغليف مطفي، Spot UV خفيف على العنوان، وبروز خفيف اختياري.",
        "التجليد المقترح: خياطة مع غلاف ورقي للمتانة، على أن تؤكد المطبعة الملاءمة بعد معرفة نوع الورق والسمك.",
        "عرض الكعب لا يُحسب قبل اعتماد الورق والسماكة والعدد النهائي للصفحات.",
        "الملف النهائي للطباعة المعتمد لدى المطبعة يجب أن يكون CMYK وPDF/X-4 مع تقرير Preflight.",
    ]
    for x in items:
        p = spec.add_paragraph(x, style="Body Text Arabic")
        set_para_rtl(p, WD_ALIGN_PARAGRAPH.JUSTIFY)
    p = spec.add_paragraph("بيانات ما تزال مطلوبة قبل الطباعة التجارية", style="Main Heading")
    for x in ["اسم المؤلف وسيرته المختصرة.", "اسم الناشر وشعاره وبيانات التواصل.", "ISBN والباركود والسعر والتصنيف.", "رابط QR عند الحاجة.", "نوع الورق ووزنه وسماكته ومواصفات التجليد من المطبعة."]:
        p = spec.add_paragraph(x, style="Numbered List Arabic")
        set_para_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
    spec.save(SPEC_DOCX)


def run_libreoffice_update_and_export():
    script = ROOT / "lo_update.py"
    script.write_text(r'''import pathlib, subprocess, time, uno
from com.sun.star.beans import PropertyValue

def prop(name, value):
    p=PropertyValue(); p.Name=name; p.Value=value; return p

proc=subprocess.Popen(['soffice','--headless','--nologo','--nodefault','--nofirststartwizard','--norestore','--accept=socket,host=127.0.0.1,port=2002;urp;StarOffice.ComponentContext'])
try:
    local=uno.getComponentContext()
    resolver=local.ServiceManager.createInstanceWithContext('com.sun.star.bridge.UnoUrlResolver',local)
    ctx=None
    for _ in range(60):
        try:
            ctx=resolver.resolve('uno:socket,host=127.0.0.1,port=2002;urp;StarOffice.ComponentContext'); break
        except Exception: time.sleep(1)
    if ctx is None: raise RuntimeError('LibreOffice connection failed')
    desktop=ctx.ServiceManager.createInstanceWithContext('com.sun.star.frame.Desktop',ctx)
    src=pathlib.Path('output/دوائر_الإنجاز_الحكومي_النسخة_النهائية.docx').resolve().as_uri()
    doc=desktop.loadComponentFromURL(src,'_blank',0,(prop('Hidden',True),prop('UpdateDocMode',3)))
    if doc is None: raise RuntimeError('Could not load DOCX')
    try: doc.getTextFields().refresh()
    except Exception: pass
    try:
        idx=doc.getDocumentIndexes()
        for i in range(idx.getCount()): idx.getByIndex(i).update()
        print('indexes',idx.getCount())
    except Exception as e: print('index update warning',e)
    try: doc.calculateAll()
    except Exception: pass
    time.sleep(5)
    out_docx=pathlib.Path('output/دوائر_الإنجاز_الحكومي_النسخة_النهائية.docx').resolve().as_uri()
    out_pdf=pathlib.Path('output/دوائر_الإنجاز_الحكومي_النسخة_النهائية.pdf').resolve().as_uri()
    doc.storeAsURL(out_docx,(prop('FilterName','Office Open XML Text'),prop('Overwrite',True)))
    try:
        idx=doc.getDocumentIndexes()
        for i in range(idx.getCount()): idx.getByIndex(i).update()
    except Exception: pass
    time.sleep(3)
    doc.storeToURL(out_pdf,(prop('FilterName','writer_pdf_Export'),prop('Overwrite',True),prop('UseTaggedPDF',True),prop('ExportBookmarks',True),prop('ExportLinksRelativeFsys',False)))
    doc.close(True)
finally:
    proc.terminate()
''', encoding="utf-8")
    subprocess.run(["/usr/bin/python3", str(script)], check=True)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(OUT), str(SPEC_DOCX)], check=True)
    generated = OUT / (SPEC_DOCX.stem + ".pdf")
    if generated.exists() and generated != SPEC_PDF:
        generated.replace(SPEC_PDF)


def audit_pdf(build_info: dict):
    pdf = fitz.open(PDF_OUT)
    page_count = len(pdf)
    expected_w = 17 / 2.54 * 72
    expected_h = 24 / 2.54 * 72
    wrong_size = []
    out_of_bounds = []
    blank_pages = []
    fonts = set()
    page_images = []
    for i, page in enumerate(pdf):
        r = page.rect
        if abs(r.width - expected_w) > 2 or abs(r.height - expected_h) > 2:
            wrong_size.append(i + 1)
        text = page.get_text("text").strip()
        if not text:
            blank_pages.append(i + 1)
        for block in page.get_text("blocks"):
            x0, y0, x1, y1 = block[:4]
            if x0 < -1 or y0 < -1 or x1 > r.width + 1 or y1 > r.height + 1:
                out_of_bounds.append((i + 1, [x0, y0, x1, y1]))
        for f in page.get_fonts(full=True):
            fonts.add(f[3])
        pix = page.get_pixmap(matrix=fitz.Matrix(1.2, 1.2), alpha=False)
        img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
        pth = OUT / f"page_{i+1:03d}.png"
        img.save(pth)
        page_images.append(pth)
    # contact sheets 12 pages per sheet
    per = 12
    for start in range(0, len(page_images), per):
        imgs = [Image.open(x) for x in page_images[start:start+per]]
        thumb_w = 300
        thumbs = []
        for im in imgs:
            ratio = thumb_w / im.width
            thumbs.append(im.resize((thumb_w, int(im.height * ratio))))
        rows = 4
        cols = 3
        th = max(x.height for x in thumbs)
        sheet = Image.new("RGB", (cols*thumb_w, rows*th), "white")
        d = ImageDraw.Draw(sheet)
        for j, im in enumerate(thumbs):
            x = (j % cols) * thumb_w
            y = (j // cols) * th
            sheet.paste(im, (x, y))
            d.text((x+5, y+5), str(start+j+1), fill="black")
        sheet.save(CONTACT_DIR / f"contact_{start//per+1:03d}.jpg", quality=88)
    final_doc = Document(DOCX_OUT)
    style_counts = {}
    for p in final_doc.paragraphs:
        name = p.style.name if p.style else ""
        style_counts[name] = style_counts.get(name, 0) + 1
    all_text = "\n".join(clean_text(p.text) for p in final_doc.paragraphs if clean_text(p.text))
    cited = sorted({int(x) for x in re.findall(r"\[(\d{1,3})\]", all_text)})
    refs = sorted({int(x) for x in re.findall(r"^\[(\d{1,3})\]", all_text, flags=re.M)})
    missing_refs = sorted(set(cited) - set(refs))
    unused_refs = sorted(set(refs) - set(cited))
    audit = {
        **build_info,
        "pdf_pages": page_count,
        "page_size_ok": not wrong_size,
        "wrong_size_pages": wrong_size,
        "out_of_bounds_blocks": out_of_bounds[:20],
        "blank_pages": blank_pages,
        "pdf_fonts": sorted(fonts),
        "style_counts": style_counts,
        "citations": cited,
        "references": refs,
        "missing_reference_numbers": missing_refs,
        "unused_reference_numbers": unused_refs,
        "contact_sheets": len(list(CONTACT_DIR.glob('*.jpg'))),
    }
    AUDIT_JSON.write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8")
    qa = Document()
    create_styles(qa)
    set_page_geometry(qa.sections[0])
    p = qa.add_paragraph("تقرير التدقيق النهائي", style="Book Title")
    set_para_rtl(p, WD_ALIGN_PARAGRAPH.CENTER)
    lines = [
        f"عدد صفحات PDF النهائي: {page_count} صفحة.",
        f"عدد الأجزاء المنسقة: {build_info['styled']['parts']}.",
        f"عدد الفصول المنسقة: {build_info['styled']['chapters']}.",
        f"عدد الجداول الفعلية: {build_info['final_tables']}، وعدد عناوين الجداول: {build_info['styled']['tables_captions']}.",
        f"عدد عناوين الأشكال: {build_info['styled']['figure_captions']}.",
        f"عدد المقاطع: {build_info['final_sections']}.",
        f"عدد الروابط القابلة للنقر التي أُنشئت: {build_info['hyperlinks_created']}.",
        f"مقاس الصفحات 17 × 24 سم: {'مطابق' if not wrong_size else 'توجد صفحات غير مطابقة'}.",
        f"النصوص خارج حدود الصفحة: {'لا يوجد' if not out_of_bounds else 'توجد حالات تحتاج فحصًا'}.",
        f"الإحالات دون مرجع مقابل: {missing_refs if missing_refs else 'لا يوجد'}.",
        f"المراجع غير المستخدمة: {unused_refs if unused_refs else 'لا يوجد'}.",
        f"الصفحات البيضاء الناتجة عن بدايات الصفحات الفردية: {blank_pages if blank_pages else 'لا يوجد'}.",
        f"عدد لوحات المراجعة المصغرة: {len(list(CONTACT_DIR.glob('*.jpg')))}.",
    ]
    for line in lines:
        p = qa.add_paragraph(line, style="Body Text Arabic")
        set_para_rtl(p, WD_ALIGN_PARAGRAPH.JUSTIFY)
    p = qa.add_paragraph("حدود الاعتماد الطباعي", style="Main Heading")
    for line in [
        "ملف PDF المرفق نسخة مراجعة متطابقة مع Word، لكنه ليس شهادة PDF/X-4 من مطبعة.",
        "عرض الكعب وملف الغلاف الكامل ينتظران مواصفات الورق والتجليد وبيانات النشر التجارية.",
        "الملف يحافظ على محتوى المخطوط؛ لم تُخترع بيانات مؤلف أو ناشر أو ISBN أو سعر.",
    ]:
        p = qa.add_paragraph(line, style="Body Text Arabic")
        set_para_rtl(p, WD_ALIGN_PARAGRAPH.JUSTIFY)
    qa.save(QA_DOCX)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(OUT), str(QA_DOCX)], check=True)
    generated = OUT / (QA_DOCX.stem + ".pdf")
    if generated.exists() and generated != QA_PDF:
        generated.replace(QA_PDF)
    return audit


def make_readme(audit: dict):
    README.write_text(
        """دوائر الإنجاز الحكومي — حزمة التسليم النهائية\n\n"
        "تتضمن الحزمة:\n"
        "1. النسخة النهائية القابلة للتحرير بصيغة Word.\n"
        "2. نسخة PDF مطابقة للمراجعة.\n"
        "3. الغلاف الأمامي والخلفي بدقة 300 DPI.\n"
        "4. دليل كعب غير مقياس حتى وصول مواصفات المطبعة.\n"
        "5. مواصفات الطباعة والتجليد.\n"
        "6. تقرير التدقيق وملف audit.json.\n"
        "7. لوحات مراجعة مصغرة لجميع الصفحات.\n\n"
        f"عدد صفحات النسخة النهائية: {audit['pdf_pages']}\n"
        "مهم: لا يزال ملف الغلاف الكامل CMYK/PDF-X وعرض الكعب النهائي مرتبطين بمواصفات الورق والتجليد وبيانات المؤلف والناشر وISBN.\n",
        encoding="utf-8",
    )


def package():
    package_dir = ROOT / "package"
    if package_dir.exists():
        shutil.rmtree(package_dir)
    package_dir.mkdir()
    for p in [DOCX_OUT, PDF_OUT, FRONT_COVER, BACK_COVER, SPINE_GUIDE, SPEC_DOCX, SPEC_PDF, QA_DOCX, QA_PDF, AUDIT_JSON, README]:
        if p.exists():
            shutil.copy2(p, package_dir / p.name)
    shutil.copytree(CONTACT_DIR, package_dir / "لوحات_المراجعة")
    zip_path = ROOT / "دوائر_الإنجاز_الحكومي_الحزمة_النهائية.zip"
    if zip_path.exists():
        zip_path.unlink()
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as z:
        for p in package_dir.rglob("*"):
            if p.is_file():
                z.write(p, p.relative_to(package_dir))
    return zip_path


def main():
    build_info = build_book()
    make_covers()
    make_support_docs()
    run_libreoffice_update_and_export()
    audit = audit_pdf(build_info)
    make_readme(audit)
    package()
    print(json.dumps(audit, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
