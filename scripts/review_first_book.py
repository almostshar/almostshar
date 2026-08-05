#!/usr/bin/env python3
from __future__ import annotations

import json
import math
import os
import re
import shutil
import subprocess
import time
import zipfile
from collections import Counter
from pathlib import Path
from typing import Iterable, Optional

import fitz
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path.cwd()
SRC = ROOT / "source.docx"
OUT = ROOT / "output"
OUT.mkdir(exist_ok=True)
WORK = OUT / "الحوكمة_من_الصفر_قبل_تحديث_الفهارس.docx"
FINAL = OUT / "الحوكمة_من_الصفر_النسخة_النهائية_المراجعة.docx"
PDF = OUT / "الحوكمة_من_الصفر_النسخة_النهائية_المراجعة.pdf"
AUDIT = OUT / "audit.json"
QA_DOCX = OUT / "مراجعة_بصرية_الحوكمة_من_الصفر.docx"
CONTACT_DIR = OUT / "contact_sheets"
CONTACT_DIR.mkdir(exist_ok=True)
DOCX_PACKAGE = ROOT / "docx_package"
QA_PACKAGE = ROOT / "qa_docx_package"

BOOK_TITLE = "الحوكمة من الصفر"
PETROL = "1D4E5F"
LIGHT_GRAY = "F2F2F2"
WHITE = "FFFFFF"
MID_GRAY = "808080"
BLACK = "000000"

ARABIC_RE = re.compile(r"[\u0600-\u06FF]")
CHAPTER_RE = re.compile(r"^الفصل\s+(?:الأول|الثاني|الثالث|الرابع|الخامس|السادس|السابع|الثامن|التاسع|العاشر|الحادي\s+عشر|الثاني\s+عشر|الثالث\s+عشر|الرابع\s+عشر|الخامس\s+عشر|[0-9٠-٩]+)")


def clean_text(value: str) -> str:
    return re.sub(r"\s+", " ", (value or "").replace("\u200f", "").replace("\u200e", "")).strip()


def arabic_ratio(value: str) -> float:
    letters = re.findall(r"[A-Za-z\u0600-\u06FF]", value or "")
    if not letters:
        return 0.0
    return sum(1 for c in letters if "\u0600" <= c <= "\u06FF") / len(letters)


def ensure(parent, tag: str):
    child = parent.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        parent.append(child)
    return child


def set_font_xml(rpr, font_name: str, size_pt: float, bold: Optional[bool] = None, color: Optional[str] = None, rtl: bool = True):
    fonts = ensure(rpr, "w:rFonts")
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), font_name)
    sz = ensure(rpr, "w:sz")
    sz.set(qn("w:val"), str(round(size_pt * 2)))
    szcs = ensure(rpr, "w:szCs")
    szcs.set(qn("w:val"), str(round(size_pt * 2)))
    if bold is not None:
        for name in ("w:b", "w:bCs"):
            el = ensure(rpr, name)
            el.set(qn("w:val"), "1" if bold else "0")
    if color:
        ensure(rpr, "w:color").set(qn("w:val"), color)
    if rtl:
        ensure(rpr, "w:rtl").set(qn("w:val"), "1")
    lang = ensure(rpr, "w:lang")
    lang.set(qn("w:val"), "ar-SA" if rtl else "en-US")
    lang.set(qn("w:eastAsia"), "ar-SA" if rtl else "en-US")
    lang.set(qn("w:bidi"), "ar-SA")


def set_run_font(run, font_name: str, size_pt: float, bold: Optional[bool] = None, color: Optional[str] = None, rtl: bool = True):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    set_font_xml(run._r.get_or_add_rPr(), font_name, size_pt, bold, color, rtl)


def set_para_bidi(p, alignment: Optional[int] = None):
    ppr = p._p.get_or_add_pPr()
    ensure(ppr, "w:bidi").set(qn("w:val"), "1")
    if alignment is not None:
        p.alignment = alignment
        val = {
            WD_ALIGN_PARAGRAPH.LEFT: "left",
            WD_ALIGN_PARAGRAPH.RIGHT: "right",
            WD_ALIGN_PARAGRAPH.CENTER: "center",
            WD_ALIGN_PARAGRAPH.JUSTIFY: "both",
        }.get(alignment, "right")
        ensure(ppr, "w:jc").set(qn("w:val"), val)


def set_style(doc: Document, name: str, font_name: str, size: float, bold=False,
              alignment=WD_ALIGN_PARAGRAPH.RIGHT, before=0, after=5, line=1.2,
              first_indent: Optional[float] = None, left_indent: Optional[float] = None,
              right_indent: Optional[float] = None, keep_next=False, keep_lines=True,
              outline: Optional[int] = None, color=BLACK):
    styles = doc.styles
    try:
        style = styles[name]
        if style.type != WD_STYLE_TYPE.PARAGRAPH:
            style.element.getparent().remove(style.element)
            style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    except KeyError:
        style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = font_name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    set_font_xml(style.element.get_or_add_rPr(), font_name, size, bold, color, True)
    pf = style.paragraph_format
    pf.alignment = alignment
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.first_line_indent = Cm(first_indent) if first_indent is not None else None
    pf.left_indent = Cm(left_indent) if left_indent is not None else None
    pf.right_indent = Cm(right_indent) if right_indent is not None else None
    pf.keep_with_next = keep_next
    pf.keep_together = keep_lines
    pf.widow_control = True
    ppr = style.element.get_or_add_pPr()
    ensure(ppr, "w:bidi").set(qn("w:val"), "1")
    ensure(ppr, "w:jc").set(qn("w:val"), {
        WD_ALIGN_PARAGRAPH.LEFT: "left",
        WD_ALIGN_PARAGRAPH.RIGHT: "right",
        WD_ALIGN_PARAGRAPH.CENTER: "center",
        WD_ALIGN_PARAGRAPH.JUSTIFY: "both",
    }.get(alignment, "right"))
    if outline is not None:
        ensure(ppr, "w:outlineLvl").set(qn("w:val"), str(outline))
    return style


def create_or_update_styles(doc: Document):
    # Core body and hierarchy
    set_style(doc, "Normal", "Amiri", 14, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 5, 1.2, .6)
    set_style(doc, "Body Text", "Amiri", 14, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 5, 1.2, .6)
    set_style(doc, "First Paragraph after Heading", "Amiri", 14, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 5, 1.2, 0)
    set_style(doc, "Chapter Title", "Noto Kufi Arabic", 22, True, WD_ALIGN_PARAGRAPH.RIGHT, 56, 20, 1.08, 0, keep_next=True, outline=1, color=PETROL)
    set_style(doc, "Main Heading", "Noto Kufi Arabic", 17, True, WD_ALIGN_PARAGRAPH.RIGHT, 14, 7, 1.08, 0, keep_next=True, outline=2, color=PETROL)
    set_style(doc, "Subheading", "Noto Kufi Arabic", 15, True, WD_ALIGN_PARAGRAPH.RIGHT, 10, 5, 1.08, 0, keep_next=True, outline=3, color=PETROL)
    set_style(doc, "Unnumbered Internal Heading", "Noto Kufi Arabic", 15, True, WD_ALIGN_PARAGRAPH.RIGHT, 10, 5, 1.08, 0, keep_next=True, outline=3, color=PETROL)
    set_style(doc, "Numbered List", "Amiri", 13.5, False, WD_ALIGN_PARAGRAPH.RIGHT, 0, 4, 1.18, 0, right_indent=.7)
    set_style(doc, "Long Quotation", "Amiri", 13, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 8, 8, 1.15, 0, left_indent=1, right_indent=1)
    # Religious styles were prepared even though no genuine passages were found.
    set_style(doc, "Qur’anic Verse", "Amiri Quran", 15, False, WD_ALIGN_PARAGRAPH.CENTER, 8, 8, 1.15, 0)
    set_style(doc, "Prophetic Hadith", "Amiri", 13.5, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 6, 6, 1.15, 0, left_indent=.8, right_indent=.8)
    # Tables, figures, boxes, references
    set_style(doc, "Table Titles", "Amiri", 10.5, True, WD_ALIGN_PARAGRAPH.RIGHT, 8, 4, 1.05, 0, keep_next=True)
    set_style(doc, "Table Title", "Amiri", 10.5, True, WD_ALIGN_PARAGRAPH.RIGHT, 8, 4, 1.05, 0, keep_next=True)
    set_style(doc, "Table Text", "Amiri", 11.5, False, WD_ALIGN_PARAGRAPH.RIGHT, 0, 2, 1.05, 0)
    set_style(doc, "Table Source", "Amiri", 10.5, False, WD_ALIGN_PARAGRAPH.RIGHT, 3, 7, 1.0, 0)
    set_style(doc, "Figure Title", "Amiri", 10.5, True, WD_ALIGN_PARAGRAPH.CENTER, 5, 3, 1.05, 0, keep_next=True)
    set_style(doc, "Figure Caption", "Amiri", 10.5, False, WD_ALIGN_PARAGRAPH.CENTER, 3, 7, 1.05, 0)
    set_style(doc, "Box Title", "Noto Kufi Arabic", 13, True, WD_ALIGN_PARAGRAPH.RIGHT, 0, 3, 1.08, 0, keep_next=True, color=PETROL)
    set_style(doc, "Box Text", "Amiri", 12.5, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 0, 1.2, 0)
    set_style(doc, "References", "Amiri", 11.5, False, WD_ALIGN_PARAGRAPH.RIGHT, 0, 4, 1.15, -0.7, right_indent=.7)
    set_style(doc, "Footnote Text", "Amiri", 10.5, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 2, 1.0, 0)
    set_style(doc, "Endnote Text", "Amiri", 10.5, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 2, 1.0, 0)
    # Built-in headings used by automatic indexes are kept RTL.
    for name, size, outline in (("Heading 1", 22, 0), ("Heading 2", 18, 1), ("Heading 3", 15, 2), ("Heading 4", 14, 3)):
        set_style(doc, name, "Noto Kufi Arabic", size, True, WD_ALIGN_PARAGRAPH.RIGHT, 10, 6, 1.08, 0, keep_next=True, outline=outline, color=PETROL)
    for name in ("TOC 1", "TOC 2", "TOC 3", "TOC 4", "Table of Figures"):
        try:
            set_style(doc, name, "Amiri", 11.5, False, WD_ALIGN_PARAGRAPH.RIGHT, 0, 2, 1.05, 0)
        except Exception:
            pass


def set_document_defaults(doc: Document):
    styles_root = doc.styles.element
    doc_defaults = styles_root.find(qn("w:docDefaults"))
    if doc_defaults is None:
        doc_defaults = OxmlElement("w:docDefaults")
        styles_root.insert(0, doc_defaults)
    rpr_default = ensure(doc_defaults, "w:rPrDefault")
    rpr = ensure(rpr_default, "w:rPr")
    set_font_xml(rpr, "Amiri", 14, False, BLACK, True)
    ppr_default = ensure(doc_defaults, "w:pPrDefault")
    ppr = ensure(ppr_default, "w:pPr")
    ensure(ppr, "w:bidi").set(qn("w:val"), "1")
    settings = doc.settings.element
    ensure(settings, "w:mirrorMargins")
    ensure(settings, "w:evenAndOddHeaders")
    ensure(settings, "w:updateFields").set(qn("w:val"), "true")
    theme_lang = ensure(settings, "w:themeFontLang")
    theme_lang.set(qn("w:val"), "ar-SA")
    theme_lang.set(qn("w:eastAsia"), "ar-SA")
    theme_lang.set(qn("w:bidi"), "ar-SA")
    ensure(settings, "w:doNotAutoHyphenate")


def section_geometry(section):
    section.page_width = Cm(17)
    section.page_height = Cm(24)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.4)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.8)
    section.gutter = Cm(.7)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)
    sectpr = section._sectPr
    pgmar = sectpr.find(qn("w:pgMar"))
    if pgmar is not None:
        pgmar.set(qn("w:gutter"), "397")


def all_story_paragraphs(doc: Document):
    yielded = set()
    for p in doc.paragraphs:
        yielded.add(id(p._p))
        yield p
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if id(p._p) not in yielded:
                        yielded.add(id(p._p))
                        yield p
    for sec in doc.sections:
        for story in (sec.header, sec.even_page_header, sec.first_page_header,
                      sec.footer, sec.even_page_footer, sec.first_page_footer):
            for p in story.paragraphs:
                if id(p._p) not in yielded:
                    yielded.add(id(p._p))
                    yield p
            for table in story.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            if id(p._p) not in yielded:
                                yielded.add(id(p._p))
                                yield p


def style_name(p) -> str:
    try:
        return p.style.name or ""
    except Exception:
        return ""


def format_paragraph(p):
    text = clean_text(p.text)
    st = style_name(p)
    lower = st.lower()
    ratio = arabic_ratio(text)
    center_styles = {"Figure Title", "Figure Caption", "Qur’anic Verse", "Title", "Subtitle", "Book Title", "Half Title"}
    justify_styles = {"Body Text", "First Paragraph after Heading", "Long Quotation", "Prophetic Hadith", "Box Text", "Footnote Text", "Endnote Text"}
    right_styles = {"Chapter Title", "Main Heading", "Subheading", "Unnumbered Internal Heading", "Numbered List", "Table Titles", "Table Title", "Table Text", "Table Source", "Box Title", "References"}
    if st in center_styles or "title page" in lower:
        align = WD_ALIGN_PARAGRAPH.CENTER
    elif st in justify_styles or (ratio >= .35 and len(text) > 120 and st not in right_styles):
        align = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif st == "References" and ratio < .35 and text:
        align = WD_ALIGN_PARAGRAPH.LEFT
    elif st.startswith("TOC") or st == "Table of Figures":
        align = WD_ALIGN_PARAGRAPH.RIGHT
    elif st in right_styles or ratio >= .15:
        align = WD_ALIGN_PARAGRAPH.RIGHT
    else:
        align = p.alignment
    set_para_bidi(p, align)
    pf = p.paragraph_format
    pf.widow_control = True
    if st in {"Chapter Title", "Main Heading", "Subheading", "Unnumbered Internal Heading", "Table Titles", "Table Title", "Figure Title", "Box Title"} or st.startswith("Heading"):
        pf.keep_with_next = True
        pf.keep_together = True
    else:
        pf.keep_together = True
    if st == "Body Text":
        pf.first_line_indent = Cm(.6)
        pf.space_before = Pt(0)
        pf.space_after = Pt(5)
        pf.line_spacing = 1.2
    elif st == "First Paragraph after Heading":
        pf.first_line_indent = Cm(0)
        pf.space_before = Pt(0)
        pf.space_after = Pt(5)
        pf.line_spacing = 1.2
    # Font enforcement is style-aware and preserves bold/italics for inline emphasis.
    heading_sizes = {
        "Chapter Title": 22,
        "Main Heading": 17,
        "Subheading": 15,
        "Unnumbered Internal Heading": 15,
        "Heading 1": 22,
        "Heading 2": 18,
        "Heading 3": 15,
        "Heading 4": 14,
        "Box Title": 13,
    }
    body_sizes = {
        "Body Text": 14,
        "First Paragraph after Heading": 14,
        "Numbered List": 13.5,
        "Long Quotation": 13,
        "Qur’anic Verse": 15,
        "Prophetic Hadith": 13.5,
        "Table Titles": 10.5,
        "Table Title": 10.5,
        "Table Text": 11.5,
        "Table Source": 10.5,
        "Figure Title": 10.5,
        "Figure Caption": 10.5,
        "Box Text": 12.5,
        "References": 11.5,
        "Footnote Text": 10.5,
        "Endnote Text": 10.5,
    }
    for run in p.runs:
        run_ratio = arabic_ratio(run.text)
        rtl = run_ratio >= .10 or ratio >= .15
        if st in heading_sizes:
            set_run_font(run, "Noto Kufi Arabic", heading_sizes[st], run.bold if run.bold is not None else True, PETROL if st not in {"Box Title"} else PETROL, True)
        else:
            size = body_sizes.get(st, 14 if ratio >= .15 else (run.font.size.pt if run.font.size else 11.5))
            font = "Amiri Quran" if st == "Qur’anic Verse" else "Amiri"
            set_run_font(run, font, size, run.bold, None, rtl)


def set_cell_shading(cell, fill: str):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = ensure(tcpr, "w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tcpr = cell._tc.get_or_add_tcPr()
    mar = ensure(tcpr, "w:tcMar")
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        el = ensure(mar, f"w:{side}")
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")


def set_table_borders(tblpr, color=MID_GRAY, size="4"):
    borders = ensure(tblpr, "w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = ensure(borders, f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def is_box_table(table) -> bool:
    if len(table.rows) == 1 and len(table.columns) == 1:
        styles = [style_name(p) for p in table.cell(0, 0).paragraphs]
        if "Box Title" in styles or "Box Text" in styles:
            return True
    for row in table.rows:
        for cell in row.cells:
            if any(style_name(p) in {"Box Title", "Box Text"} for p in cell.paragraphs):
                return True
    return False


def format_box_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    tblpr = table._tbl.tblPr
    ensure(tblpr, "w:bidiVisual")
    borders = ensure(tblpr, "w:tblBorders")
    for edge in ("top", "left", "bottom", "insideH", "insideV"):
        el = ensure(borders, f"w:{edge}")
        el.set(qn("w:val"), "nil")
    right = ensure(borders, "w:right")
    right.set(qn("w:val"), "single")
    right.set(qn("w:sz"), "18")
    right.set(qn("w:space"), "0")
    right.set(qn("w:color"), PETROL)
    for row in table.rows:
        trpr = row._tr.get_or_add_trPr()
        ensure(trpr, "w:cantSplit")
        for cell in row.cells:
            set_cell_shading(cell, LIGHT_GRAY)
            set_cell_margins(cell, 227, 227, 227, 227)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                format_paragraph(p)


def format_content_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    tblpr = table._tbl.tblPr
    ensure(tblpr, "w:bidiVisual")
    set_table_borders(tblpr)
    for ri, row in enumerate(table.rows):
        trpr = row._tr.get_or_add_trPr()
        ensure(trpr, "w:cantSplit")
        if ri == 0:
            ensure(trpr, "w:tblHeader")
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_shading(cell, PETROL if ri == 0 else (LIGHT_GRAY if ri % 2 == 0 else WHITE))
            for p in cell.paragraphs:
                p.style = "Table Text"
                set_para_bidi(p, WD_ALIGN_PARAGRAPH.RIGHT)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.05
                p.paragraph_format.keep_together = True
                for run in p.runs:
                    set_run_font(run, "Amiri", 11.5, True if ri == 0 else run.bold, WHITE if ri == 0 else BLACK, True)


def format_tables(doc: Document):
    boxes = 0
    contents = 0
    for table in doc.tables:
        if is_box_table(table):
            boxes += 1
            format_box_table(table)
        else:
            contents += 1
            format_content_table(table)
    return boxes, contents


def format_headers_footers(doc: Document):
    settings = doc.settings.element
    ensure(settings, "w:evenAndOddHeaders")
    for i, sec in enumerate(doc.sections):
        section_geometry(sec)
        if i >= 2:
            sec.start_type = WD_SECTION.ODD_PAGE
        for story_name, story, align in (
            ("odd_header", sec.header, WD_ALIGN_PARAGRAPH.RIGHT),
            ("even_header", sec.even_page_header, WD_ALIGN_PARAGRAPH.LEFT),
            ("first_header", sec.first_page_header, WD_ALIGN_PARAGRAPH.RIGHT),
            ("odd_footer", sec.footer, WD_ALIGN_PARAGRAPH.RIGHT),
            ("even_footer", sec.even_page_footer, WD_ALIGN_PARAGRAPH.LEFT),
            ("first_footer", sec.first_page_footer, WD_ALIGN_PARAGRAPH.RIGHT),
        ):
            for p in story.paragraphs:
                set_para_bidi(p, align)
                for run in p.runs:
                    set_run_font(run, "Amiri", 10.5, run.bold, None, True)


def format_note_xml(docx_path: Path):
    tmp = OUT / "note_patch"
    if tmp.exists():
        shutil.rmtree(tmp)
    tmp.mkdir()
    with zipfile.ZipFile(docx_path) as z:
        z.extractall(tmp)
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    from lxml import etree
    changed = False
    for rel in ("word/footnotes.xml", "word/endnotes.xml"):
        path = tmp / rel
        if not path.exists():
            continue
        tree = etree.parse(str(path))
        for p in tree.xpath("//w:p", namespaces=ns):
            ppr = p.find(qn("w:pPr"))
            if ppr is None:
                ppr = etree.Element(qn("w:pPr"))
                p.insert(0, ppr)
            if ppr.find(qn("w:bidi")) is None:
                ppr.insert(0, etree.Element(qn("w:bidi")))
            jc = ppr.find(qn("w:jc"))
            if jc is None:
                jc = etree.SubElement(ppr, qn("w:jc"))
            jc.set(qn("w:val"), "right")
            for r in p.findall(".//" + qn("w:r")):
                rpr = r.find(qn("w:rPr"))
                if rpr is None:
                    rpr = etree.Element(qn("w:rPr"))
                    r.insert(0, rpr)
                if rpr.find(qn("w:rtl")) is None:
                    etree.SubElement(rpr, qn("w:rtl"))
                fonts = rpr.find(qn("w:rFonts"))
                if fonts is None:
                    fonts = etree.SubElement(rpr, qn("w:rFonts"))
                for key in ("ascii", "hAnsi", "eastAsia", "cs"):
                    fonts.set(qn(f"w:{key}"), "Amiri")
                lang = rpr.find(qn("w:lang"))
                if lang is None:
                    lang = etree.SubElement(rpr, qn("w:lang"))
                lang.set(qn("w:val"), "ar-SA")
                lang.set(qn("w:bidi"), "ar-SA")
        tree.write(str(path), xml_declaration=True, encoding="UTF-8", standalone="yes")
        changed = True
    if changed:
        patched = OUT / "notes_patched.docx"
        with zipfile.ZipFile(patched, "w", zipfile.ZIP_DEFLATED) as z:
            for p in tmp.rglob("*"):
                if p.is_file():
                    z.write(p, p.relative_to(tmp))
        shutil.copy2(patched, docx_path)
    shutil.rmtree(tmp)


def package_text(path: Path) -> str:
    from lxml import etree
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml")
    root = etree.fromstring(xml)
    texts = root.xpath("//w:body//w:t/text()", namespaces={"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"})
    return "".join(texts)


def apply_all_formatting():
    if not SRC.exists():
        raise FileNotFoundError(SRC)
    source_text = package_text(SRC)
    doc = Document(SRC)
    create_or_update_styles(doc)
    set_document_defaults(doc)
    for sec in doc.sections:
        section_geometry(sec)
    for p in all_story_paragraphs(doc):
        format_paragraph(p)
    box_count, content_table_count = format_tables(doc)
    format_headers_footers(doc)
    cp = doc.core_properties
    cp.title = BOOK_TITLE
    cp.subject = "دليل عملي للشركات السعودية من التأسيس إلى النضج"
    cp.language = "ar-SA"
    cp.comments = "نسخة مراجعة نهائية منسقة للطباعة باتجاه عربي من اليمين إلى اليسار"
    doc.save(WORK)
    format_note_xml(WORK)
    formatted_text = package_text(WORK)
    if source_text != formatted_text:
        raise RuntimeError("Visible document text changed during formatting pass")
    return {
        "source_text_characters": len(source_text),
        "text_preserved_before_index_refresh": True,
        "box_tables": box_count,
        "content_tables": content_table_count,
    }


def run_libreoffice():
    script = ROOT / "lo_update_first.py"
    script.write_text(r'''import pathlib, subprocess, time, uno
from com.sun.star.beans import PropertyValue

def prop(name, value):
    p=PropertyValue(); p.Name=name; p.Value=value; return p

proc=subprocess.Popen(['soffice','--headless','--nologo','--nodefault','--nofirststartwizard','--norestore','--accept=socket,host=127.0.0.1,port=2002;urp;StarOffice.ComponentContext'])
try:
    local=uno.getComponentContext()
    resolver=local.ServiceManager.createInstanceWithContext('com.sun.star.bridge.UnoUrlResolver',local)
    ctx=None
    for _ in range(60):
        try:
            ctx=resolver.resolve('uno:socket,host=127.0.0.1,port=2002;urp;StarOffice.ComponentContext'); break
        except Exception: time.sleep(1)
    if ctx is None: raise RuntimeError('LibreOffice connection failed')
    desktop=ctx.ServiceManager.createInstanceWithContext('com.sun.star.frame.Desktop',ctx)
    src=pathlib.Path('output/الحوكمة_من_الصفر_قبل_تحديث_الفهارس.docx').resolve().as_uri()
    doc=desktop.loadComponentFromURL(src,'_blank',0,(prop('Hidden',True),prop('UpdateDocMode',3)))
    if doc is None: raise RuntimeError('Could not load DOCX')
    try: doc.getTextFields().refresh()
    except Exception: pass
    try:
        indexes=doc.getDocumentIndexes()
        for i in range(indexes.getCount()): indexes.getByIndex(i).update()
        print('indexes_updated', indexes.getCount())
    except Exception as e: print('index_warning', e)
    try: doc.calculateAll()
    except Exception: pass
    time.sleep(5)
    out_docx=pathlib.Path('output/الحوكمة_من_الصفر_النسخة_النهائية_المراجعة.docx').resolve().as_uri()
    out_pdf=pathlib.Path('output/الحوكمة_من_الصفر_النسخة_النهائية_المراجعة.pdf').resolve().as_uri()
    doc.storeAsURL(out_docx,(prop('FilterName','Office Open XML Text'),prop('Overwrite',True)))
    try:
        indexes=doc.getDocumentIndexes()
        for i in range(indexes.getCount()): indexes.getByIndex(i).update()
    except Exception: pass
    time.sleep(2)
    doc.storeToURL(out_pdf,(prop('FilterName','writer_pdf_Export'),prop('Overwrite',True),prop('UseTaggedPDF',True),prop('ExportBookmarks',True)))
    doc.close(True)
finally:
    proc.terminate()
''', encoding="utf-8")
    subprocess.run(["/usr/bin/python3", str(script)], check=True)


def inspect_docx(final_path: Path, base_info: dict):
    doc = Document(final_path)
    issues = []
    rtl_violations = []
    style_counts = Counter()
    chapter_titles = []
    caption_count = 0
    source_count = 0
    box_titles = 0
    box_texts = 0
    for p in all_story_paragraphs(doc):
        text = clean_text(p.text)
        st = style_name(p)
        style_counts[st] += 1
        if st in {"Table Titles", "Table Title"} and text.startswith("الجدول"):
            caption_count += 1
        if text.startswith("المصدر:"):
            source_count += 1
        if st == "Box Title":
            box_titles += 1
        if st == "Box Text":
            box_texts += 1
        if CHAPTER_RE.match(text):
            chapter_titles.append(text)
        ppr = p._p.pPr
        has_bidi = ppr is not None and ppr.find(qn("w:bidi")) is not None
        if text and arabic_ratio(text) >= .15 and not has_bidi:
            rtl_violations.append(text[:120])
        if text and arabic_ratio(text) >= .30 and p.alignment == WD_ALIGN_PARAGRAPH.LEFT and st != "References":
            rtl_violations.append("LEFT:" + text[:110])
    sections = len(doc.sections)
    geometry_bad = []
    for idx, sec in enumerate(doc.sections, start=1):
        vals = (sec.page_width.cm, sec.page_height.cm, sec.top_margin.cm, sec.bottom_margin.cm, sec.left_margin.cm, sec.right_margin.cm, sec.gutter.cm)
        expected = (17, 24, 2.2, 2.4, 2.0, 2.8, .7)
        if any(abs(a-b) > .04 for a, b in zip(vals, expected)):
            geometry_bad.append((idx, vals))
    boxes = sum(1 for t in doc.tables if is_box_table(t))
    contents = len(doc.tables) - boxes
    with zipfile.ZipFile(final_path) as z:
        names = z.namelist()
        document_xml = z.read("word/document.xml").decode("utf-8", errors="ignore")
        settings_xml = z.read("word/settings.xml").decode("utf-8", errors="ignore")
        bad_zip = z.testzip()
        media = [x for x in names if x.startswith("word/media/")]
        field_count = len(re.findall(r"TOC", document_xml))
        bidi_visual = document_xml.count("bidiVisual")
        revisions = len(re.findall(r"<w:(?:ins|del)\b", document_xml))
    if rtl_violations:
        issues.append(f"RTL violations: {len(rtl_violations)}")
    if geometry_bad:
        issues.append(f"Geometry violations: {len(geometry_bad)}")
    if boxes != 43:
        issues.append(f"Expected 43 box tables, found {boxes}")
    if contents != 25:
        issues.append(f"Expected 25 content tables, found {contents}")
    if caption_count != 25:
        issues.append(f"Expected 25 table captions, found {caption_count}")
    if source_count != 25:
        issues.append(f"Expected 25 table source notes, found {source_count}")
    if len(chapter_titles) != 15:
        issues.append(f"Expected 15 chapter titles, found {len(chapter_titles)}")
    if bidi_visual < len(doc.tables):
        issues.append(f"Only {bidi_visual} of {len(doc.tables)} tables have bidiVisual")
    if "mirrorMargins" not in settings_xml or "evenAndOddHeaders" not in settings_xml:
        issues.append("Mirrored margins or odd/even headers setting missing")
    if bad_zip:
        issues.append(f"Corrupt ZIP member: {bad_zip}")
    return {
        **base_info,
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "box_tables_final": boxes,
        "content_tables_final": contents,
        "sections": sections,
        "chapter_titles": chapter_titles,
        "chapter_count": len(chapter_titles),
        "table_captions": caption_count,
        "table_sources": source_count,
        "box_title_paragraphs": box_titles,
        "box_text_paragraphs": box_texts,
        "rtl_violations": rtl_violations[:30],
        "geometry_violations": geometry_bad,
        "toc_field_mentions": field_count,
        "bidi_visual_table_count": bidi_visual,
        "media_files": media,
        "tracked_revision_elements": revisions,
        "style_counts": dict(style_counts),
        "docx_zip_integrity": bad_zip is None,
        "issues": issues,
    }


def font_path(pattern: str) -> str:
    try:
        return subprocess.check_output(["fc-match", "-f", "%{file}", pattern], text=True).strip()
    except Exception:
        return "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"


def audit_pdf(info: dict):
    pdf = fitz.open(PDF)
    expected_w = 17 / 2.54 * 72
    expected_h = 24 / 2.54 * 72
    wrong_size = []
    out_bounds = []
    blank = []
    chapter_pages = {}
    orphan_heading_pages = []
    rendered = []
    for idx, page in enumerate(pdf, start=1):
        rect = page.rect
        if abs(rect.width - expected_w) > 2 or abs(rect.height - expected_h) > 2:
            wrong_size.append(idx)
        blocks = page.get_text("blocks")
        if not page.get_text("text").strip():
            blank.append(idx)
        for b in blocks:
            x0, y0, x1, y1 = b[:4]
            if x0 < -1 or y0 < -1 or x1 > rect.width + 1 or y1 > rect.height + 1:
                out_bounds.append((idx, [x0, y0, x1, y1]))
        page_dict = page.get_text("dict")
        large_spans = []
        for block in page_dict.get("blocks", []):
            for line in block.get("lines", []):
                for span in line.get("spans", []):
                    text = clean_text(span.get("text", ""))
                    size = float(span.get("size", 0))
                    if size >= 16 and text:
                        large_spans.append((text, size, span.get("bbox", [0,0,0,0])))
                        if CHAPTER_RE.match(text) and text not in chapter_pages:
                            chapter_pages[text] = idx
        if large_spans:
            bottom_large = max(large_spans, key=lambda x: x[2][3])
            if bottom_large[2][3] > rect.height * .82:
                later_body = [b for b in blocks if b[1] > bottom_large[2][3] + 2 and clean_text(str(b[4]))]
                if not later_body:
                    orphan_heading_pages.append(idx)
        pix = page.get_pixmap(matrix=fitz.Matrix(1.15, 1.15), alpha=False)
        im = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
        p = OUT / f"page_{idx:03d}.jpg"
        im.save(p, quality=88)
        rendered.append(p)
    chapter_parity_bad = {k: v for k, v in chapter_pages.items() if v % 2 == 0}
    info.update({
        "pdf_pages": len(pdf),
        "page_size_ok": not wrong_size,
        "wrong_size_pages": wrong_size,
        "out_of_bounds_blocks": out_bounds[:20],
        "blank_pages": blank,
        "chapter_opening_pages": chapter_pages,
        "chapter_openings_on_even_pdf_pages": chapter_parity_bad,
        "possible_orphan_heading_pages": orphan_heading_pages,
    })
    if wrong_size:
        info["issues"].append(f"Wrong PDF page size: {wrong_size}")
    if out_bounds:
        info["issues"].append(f"Out-of-bounds PDF blocks: {len(out_bounds)}")
    if chapter_parity_bad:
        info["issues"].append(f"Chapter openings not on recto pages: {chapter_parity_bad}")
    make_contact_sheets(rendered)
    return info


def make_contact_sheets(images: list[Path]):
    for old in CONTACT_DIR.glob("*"):
        old.unlink()
    font = ImageFont.truetype(font_path("DejaVu Sans"), 18)
    per_sheet = 15
    cols, rows = 3, 5
    thumb_w = 330
    margin = 18
    label_h = 26
    sheet_paths = []
    for start in range(0, len(images), per_sheet):
        group = images[start:start + per_sheet]
        opened = [Image.open(p).convert("RGB") for p in group]
        thumb_h = max(int(im.height * thumb_w / im.width) for im in opened)
        sheet = Image.new("RGB", (cols * thumb_w + (cols + 1) * margin, rows * (thumb_h + label_h) + (rows + 1) * margin), "white")
        draw = ImageDraw.Draw(sheet)
        for j, im in enumerate(opened):
            scale = thumb_w / im.width
            thumb = im.resize((thumb_w, int(im.height * scale)))
            x = margin + (j % cols) * (thumb_w + margin)
            y = margin + (j // cols) * (thumb_h + label_h + margin)
            sheet.paste(thumb, (x, y))
            draw.text((x + 5, y + thumb_h + 2), f"Page {start + j + 1}", fill="black", font=font)
        path = CONTACT_DIR / f"contact_{start // per_sheet + 1:02d}.jpg"
        sheet.save(path, quality=91)
        sheet_paths.append(path)
    create_qa_docx(sheet_paths)


def create_qa_docx(sheet_paths: list[Path]):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(29.7)
    sec.page_height = Cm(21)
    sec.top_margin = Cm(.6)
    sec.bottom_margin = Cm(.6)
    sec.left_margin = Cm(.6)
    sec.right_margin = Cm(.6)
    p = doc.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("الحوكمة من الصفر — لوحات المراجعة البصرية")
    set_run_font(r, "Noto Kufi Arabic", 16, True, PETROL, True)
    set_para_bidi(p, WD_ALIGN_PARAGRAPH.CENTER)
    for i, sheet in enumerate(sheet_paths):
        if i > 0 or p.text:
            p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.add_run().add_picture(str(sheet), width=Cm(28.2))
        if i < len(sheet_paths) - 1:
            p.add_run().add_break()
            p.paragraph_format.page_break_after = True
    doc.save(QA_DOCX)


def extract_package(docx: Path, dest: Path):
    if dest.exists():
        shutil.rmtree(dest)
    dest.mkdir()
    with zipfile.ZipFile(docx) as z:
        z.extractall(dest)
    # Verify that rebuilding the package is still a valid DOCX.
    test_docx = OUT / (dest.name + "_rebuilt.docx")
    with zipfile.ZipFile(test_docx, "w", zipfile.ZIP_DEFLATED) as z:
        for p in dest.rglob("*"):
            if p.is_file():
                z.write(p, p.relative_to(dest))
    Document(test_docx)


def main():
    base = apply_all_formatting()
    run_libreoffice()
    info = inspect_docx(FINAL, base)
    info = audit_pdf(info)
    AUDIT.write_text(json.dumps(info, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(info, ensure_ascii=False, indent=2))
    if info["issues"]:
        raise RuntimeError("Final audit failed: " + " | ".join(info["issues"]))
    extract_package(FINAL, DOCX_PACKAGE)
    extract_package(QA_DOCX, QA_PACKAGE)


if __name__ == "__main__":
    main()
